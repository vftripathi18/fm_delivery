import os
from datetime import datetime

import frappe
from frappe.utils import get_site_path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


# ============================================================
# FM DELIVERY EXCEL IMPORTER
# ============================================================

class ExcelImporter:

    # ============================================================
    # DAILY EXCEL HEADERS
    # ============================================================

    EXPECTED_HEADERS = [
        "Date",
        "State",
        "Hub_Name",
        "FHRID",
        "Agent_Name",
        "tripsheetId",
        "DeliveredReturn",
        "PickedForward",
        "Total"
    ]

    # ============================================================
    # INITIALIZE
    # ============================================================

    def __init__(self, import_doc):

        self.import_doc = import_doc
        self.workbook = None
        self.sheet = None

    # ============================================================
    # VALIDATE FILE
    # ============================================================

    def validate_file(self):

        if not self.import_doc.import_file:

            frappe.throw(
                "Please attach an Excel file."
            )

        if not self.import_doc.import_file.lower().endswith(".xlsx"):

            frappe.throw(
                "Only .xlsx files are allowed."
            )

    # ============================================================
    # LOAD WORKBOOK
    # ============================================================

    def load_workbook(self):

        self.validate_file()

        file_doc = frappe.get_doc(
            "File",
            {
                "file_url": self.import_doc.import_file
            }
        )

        file_path = file_doc.get_full_path()

        if not os.path.exists(file_path):

            frappe.throw(
                f"File not found: {file_path}"
            )

        self.workbook = load_workbook(
            filename=file_path,
            data_only=True,
            read_only=True
        )

        self.sheet = self.workbook.active

        return self.workbook, self.sheet

    # ============================================================
    # VALIDATE HEADERS
    # ============================================================

    def validate_headers(self):

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in self.sheet[1]
        ]

        if headers != self.EXPECTED_HEADERS:

            frappe.throw(
                f"""
                <h4>Excel Header Validation Failed</h4>

                <b>Expected:</b>
                <br>
                {'<br>'.join(self.EXPECTED_HEADERS)}

                <hr>

                <b>Found:</b>
                <br>
                {'<br>'.join(headers)}
                """
            )

    # ============================================================
    # NUMBER CONVERTER
    # ============================================================

    def to_float(self, value):

        if value in (None, "", "NULL"):

            return 0.0

        try:

            return float(value)

        except Exception:

            raise Exception(
                f"Invalid number value: {value}"
            )

    # ============================================================
    # DATE CONVERTER
    # ============================================================

    def to_date(self, value):

        if not value:

            return None

        if isinstance(value, datetime):

            return value.date()

        return value

    # ============================================================
    # FIND SALARY DECLARATION
    # ============================================================

    def get_salary_declaration(
        self,
        state,
        hub_name,
        delivery_date
    ):

        # --------------------------------------------------------
        # Find Parent Salary Declaration
        # --------------------------------------------------------

        declarations = frappe.get_all(
            "FM Salary Declaration",

            filters={
                "from_date": ["<=", delivery_date],
                "to_date": [">=", delivery_date],
            },

            fields=[
                "name",
                "from_date",
                "to_date",
            ],

            order_by="from_date desc",

            limit=2,
        )

        # --------------------------------------------------------
        # No Salary Declaration
        # --------------------------------------------------------

        if not declarations:

            raise Exception(
                "SALARY_DECLARATION_NOT_FOUND: "
                f"No Salary Declaration found for "
                f"Date '{delivery_date}'. "
                f"State '{state}', "
                f"Hub '{hub_name}'."
            )

        # --------------------------------------------------------
        # Multiple Salary Declarations
        # --------------------------------------------------------

        if len(declarations) > 1:

            raise Exception(
                "MULTIPLE_SALARY_DECLARATIONS: "
                f"Multiple Salary Declarations found "
                f"for Date '{delivery_date}'. "
                f"Please check overlapping periods."
            )

        declaration = declarations[0]

        # --------------------------------------------------------
        # Get Parent Document
        # --------------------------------------------------------

        parent_doc = frappe.get_doc(
            "FM Salary Declaration",
            declaration.name
        )

        # --------------------------------------------------------
        # Find Matching State + Hub
        # --------------------------------------------------------

        matching_rows = []

        for child in parent_doc.salary_details:

            child_state = str(
                child.state or ""
            ).strip()

            child_hub = str(
                child.hub_name or ""
            ).strip()

            if (
                child_state.lower() == state.lower()
                and
                child_hub.lower() == hub_name.lower()
            ):

                matching_rows.append(
                    child
                )

        # --------------------------------------------------------
        # Hub Not Declared
        # --------------------------------------------------------

        if not matching_rows:

            raise Exception(
                "HUB_NOT_DECLARED: "
                f"Hub '{hub_name}' is not declared "
                f"in Salary Declaration for "
                f"State '{state}' and "
                f"Date '{delivery_date}'."
            )

        # --------------------------------------------------------
        # Duplicate State + Hub
        # --------------------------------------------------------

        if len(matching_rows) > 1:

            raise Exception(
                "DUPLICATE_SALARY_CONFIGURATION: "
                f"Multiple salary configurations found "
                f"for State '{state}' and "
                f"Hub '{hub_name}'."
            )

        return (
            parent_doc,
            matching_rows[0]
        )

    # ============================================================
    # CALCULATE SALARY
    # ============================================================

    def calculate_salary(
        self,
        salary_row,
        delivered_return,
        picked_forward,
        total
    ):

        fix_salary = self.to_float(
            salary_row.fix_salary
        )

        fwd_rate = self.to_float(
            salary_row.fwd_rate
        )

        rto_rate = self.to_float(
            salary_row.rto_rate
        )

        total_rate = self.to_float(
            salary_row.total_rate
        )

        # --------------------------------------------------------
        # FIX SALARY
        # --------------------------------------------------------

        if fix_salary > 0:

            return fix_salary

        # --------------------------------------------------------
        # FWD + RTO RATE
        # --------------------------------------------------------

        if (
            fwd_rate > 0
            or
            rto_rate > 0
        ):

            return (
                picked_forward * fwd_rate
                +
                delivered_return * rto_rate
            )

        # --------------------------------------------------------
        # TOTAL RATE
        # --------------------------------------------------------

        if total_rate > 0:

            return (
                total * total_rate
            )

        # --------------------------------------------------------
        # NO VALID CONFIGURATION
        # --------------------------------------------------------

        raise Exception(
            "INVALID_SALARY_CONFIGURATION: "
            f"No valid salary configuration found "
            f"for State '{salary_row.state}', "
            f"Hub '{salary_row.hub_name}'."
        )

    # ============================================================
    # IMPORT DATA
    # ============================================================

    def import_data(self):

        attempted = 0
        imported = 0
        updated = 0
        failed = 0
        duplicates = 0

        failed_rows = []

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in self.sheet[1]
        ]

        # --------------------------------------------------------
        # Track duplicates inside the SAME Excel file
        #
        # Key:
        # Date + FHRID + TripSheet ID
        # --------------------------------------------------------

        processed_keys = set()

        # --------------------------------------------------------
        # Process Excel Rows
        # --------------------------------------------------------

        for row_no, values in enumerate(
            self.sheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            # ----------------------------------------------------
            # Skip completely empty rows
            # ----------------------------------------------------

            if not any(values):

                continue

            attempted += 1

            try:

                # ------------------------------------------------
                # Convert Row To Dictionary
                # ------------------------------------------------

                row = {
                    headers[i]:
                        values[i]
                        if i < len(values)
                        else None

                    for i in range(
                        len(headers)
                    )
                }

                # ------------------------------------------------
                # Read Daily Data
                # ------------------------------------------------

                delivery_date = self.to_date(
                    row.get("Date")
                )

                state = str(
                    row.get("State") or ""
                ).strip()

                hub_name = str(
                    row.get("Hub_Name") or ""
                ).strip()

                fhr_id = str(
                    row.get("FHRID") or ""
                ).strip()

                agent_name = str(
                    row.get("Agent_Name") or ""
                ).strip()

                tripsheet_id = str(
                    row.get("tripsheetId") or ""
                ).strip()

                delivered_return = self.to_float(
                    row.get("DeliveredReturn")
                )

                picked_forward = self.to_float(
                    row.get("PickedForward")
                )

                total = self.to_float(
                    row.get("Total")
                )

                # ------------------------------------------------
                # Required Fields
                # ------------------------------------------------

                if not delivery_date:

                    raise Exception(
                        "MISSING_DATE: Date is required."
                    )

                if not state:

                    raise Exception(
                        "MISSING_STATE: State is required."
                    )

                if not hub_name:

                    raise Exception(
                        "MISSING_HUB: Hub_Name is required."
                    )

                if not fhr_id:

                    raise Exception(
                        "MISSING_FHRID: FHRID is required."
                    )

                if not tripsheet_id:

                    raise Exception(
                        "MISSING_TRIPSHEETID: "
                        "tripsheetId is required."
                    )

                # ------------------------------------------------
                # DUPLICATE KEY
                #
                # Same:
                # Date + FHRID + TripSheet ID
                #
                # means same allocation.
                # ------------------------------------------------

                duplicate_key = (
                    str(delivery_date),
                    fhr_id.lower(),
                    tripsheet_id.lower()
                )

                # ------------------------------------------------
                # FIND SALARY DECLARATION
                # ------------------------------------------------

                (
                    salary_declaration,
                    salary_row
                ) = self.get_salary_declaration(
                    state=state,
                    hub_name=hub_name,
                    delivery_date=delivery_date
                )

                # ------------------------------------------------
                # CALCULATE SALARY
                # ------------------------------------------------

                salary = self.calculate_salary(
                    salary_row=salary_row,
                    delivered_return=delivered_return,
                    picked_forward=picked_forward,
                    total=total
                )

                # ------------------------------------------------
                # CALCULATE RCPS
                # ------------------------------------------------

                if total > 0:

                    rcps = salary / total

                else:

                    rcps = 0.0

                # =================================================
                # CHECK EXISTING DAILY RECORD
                #
                # UNIQUE LOGIC:
                #
                # Date
                # +
                # FHRID
                # +
                # TripSheet ID
                #
                # IMPORTANT:
                #
                # Same FHRID + different TripSheet
                # = DIFFERENT RECORD
                #
                # Same FHRID + same TripSheet
                # = SAME RECORD / UPDATE
                # =================================================

                existing_name = frappe.db.get_value(
                    "FM Delivery Data",
                    {
                        "date": delivery_date,
                        "fhr_id": fhr_id,
                        "tripsheet_id": tripsheet_id,
                    },
                    "name"
                )

                # ------------------------------------------------
                # UPDATE EXISTING RECORD
                # ------------------------------------------------

                if existing_name:

                    delivery_doc = frappe.get_doc(
                        "FM Delivery Data",
                        existing_name
                    )

                    delivery_doc.update({

                        "date":
                            delivery_date,

                        "state":
                            state,

                        "hub_name":
                            hub_name,

                        "fhr_id":
                            fhr_id,

                        "rider_name":
                            agent_name,

                        "tripsheet_id":
                            tripsheet_id,

                        "return_shipment":
                            delivered_return,

                        "forward_shipment":
                            picked_forward,

                        "total_shipment":
                            total,

                        "salary":
                            salary,

                        "rcps":
                            rcps,

                    })

                    delivery_doc.save(
                        ignore_permissions=True
                    )

                    updated += 1

                    # Existing DB record is not counted
                    # as an Excel duplicate.

                # ------------------------------------------------
                # CREATE NEW RECORD
                # ------------------------------------------------

                else:

                    delivery_doc = frappe.new_doc(
                        "FM Delivery Data"
                    )

                    delivery_doc.update({

                        "date":
                            delivery_date,

                        "state":
                            state,

                        "hub_name":
                            hub_name,

                        "fhr_id":
                            fhr_id,

                        "rider_name":
                            agent_name,

                        "tripsheet_id":
                            tripsheet_id,

                        "return_shipment":
                            delivered_return,

                        "forward_shipment":
                            picked_forward,

                        "total_shipment":
                            total,

                        "salary":
                            salary,

                        "rcps":
                            rcps,

                    })

                    delivery_doc.insert(
                        ignore_permissions=True
                    )

                    imported += 1

                # ------------------------------------------------
                # Mark this combination as processed
                #
                # This is useful for tracking exact duplicates
                # inside the uploaded Excel.
                # ------------------------------------------------

                if duplicate_key in processed_keys:

                    duplicates += 1

                else:

                    processed_keys.add(
                        duplicate_key
                    )

            except Exception as e:

                failed += 1

                error_message = str(e)

                # ------------------------------------------------
                # Determine Error Type
                # ------------------------------------------------

                if error_message.startswith(
                    "HUB_NOT_DECLARED:"
                ):

                    error_type = "HUB NOT DECLARED"

                elif error_message.startswith(
                    "SALARY_DECLARATION_NOT_FOUND:"
                ):

                    error_type = (
                        "SALARY DECLARATION NOT FOUND"
                    )

                elif error_message.startswith(
                    "MULTIPLE_SALARY_DECLARATIONS:"
                ):

                    error_type = (
                        "MULTIPLE SALARY DECLARATIONS"
                    )

                elif error_message.startswith(
                    "DUPLICATE_SALARY_CONFIGURATION:"
                ):

                    error_type = (
                        "DUPLICATE SALARY CONFIGURATION"
                    )

                elif error_message.startswith(
                    "INVALID_SALARY_CONFIGURATION:"
                ):

                    error_type = (
                        "INVALID SALARY CONFIGURATION"
                    )

                elif error_message.startswith(
                    "MISSING_DATE:"
                ):

                    error_type = "MISSING DATE"

                elif error_message.startswith(
                    "MISSING_STATE:"
                ):

                    error_type = "MISSING STATE"

                elif error_message.startswith(
                    "MISSING_HUB:"
                ):

                    error_type = "MISSING HUB"

                elif error_message.startswith(
                    "MISSING_FHRID:"
                ):

                    error_type = "MISSING FHRID"

                elif error_message.startswith(
                    "MISSING_TRIPSHEETID:"
                ):

                    error_type = (
                        "MISSING TRIPSHEET ID"
                    )

                elif error_message.startswith(
                    "Invalid number value:"
                ):

                    error_type = "INVALID NUMBER"

                else:

                    error_type = "IMPORT ERROR"

                # ------------------------------------------------
                # Clean Error Message
                # ------------------------------------------------

                clean_message = (
                    error_message
                    .split(":", 1)[1]
                    .strip()
                    if ":" in error_message
                    else error_message
                )

                failed_rows.append({

                    "row_no":
                        row_no,

                    "row_data":
                        values,

                    "error_type":
                        error_type,

                    "error":
                        clean_message,

                })

                # ------------------------------------------------
                # Log technical traceback
                # ------------------------------------------------

                frappe.log_error(
                    title=f"FM Delivery Import Row {row_no}",
                    message=frappe.get_traceback()
                )

                # IMPORTANT:
                # Do NOT frappe.throw here.
                # Remaining rows continue.

        # ========================================================
        # COMMIT
        # ========================================================

        frappe.db.commit()

        # ========================================================
        # ERROR REPORT
        # ========================================================

        error_file_url = None

        if failed_rows:

            error_file_url = self.create_error_report(
                failed_rows
            )

        # ========================================================
        # UPDATE IMPORT DOCUMENT
        # ========================================================

        self.import_doc.imported_rows = (
            imported + updated
        )

        self.import_doc.failed_rows = failed

        self.import_doc.status = "Completed"

        if error_file_url:

            self.import_doc.error_file = (
                error_file_url
            )

        self.import_doc.save(
            ignore_permissions=True
        )

        # ========================================================
        # RETURN SUMMARY
        # ========================================================

        return {

            "attempted":
                attempted,

            "imported":
                imported,

            "updated":
                updated,

            "failed":
                failed,

            "duplicates":
                duplicates,

            "difference":
                attempted
                -
                (
                    imported
                    +
                    updated
                    +
                    failed
                ),

            "error_file":
                error_file_url,

        }

    # ============================================================
    # CREATE ERROR REPORT
    # ============================================================

    def create_error_report(
        self,
        failed_rows
    ):

        wb = Workbook()

        ws = wb.active

        ws.title = "Errors"

        # --------------------------------------------------------
        # Error Report Headers
        # --------------------------------------------------------

        headers = [
            "Row Number"
        ] + self.EXPECTED_HEADERS + [
            "Error Type",
            "Error Message"
        ]

        ws.append(headers)

        # --------------------------------------------------------
        # Header Styling
        # --------------------------------------------------------

        bold_font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="FFFFFF"
        )

        red_fill = PatternFill(
            start_color="FF0000",
            end_color="FF0000",
            fill_type="solid"
        )

        for col_idx in range(
            1,
            len(headers) + 1
        ):

            cell = ws.cell(
                row=1,
                column=col_idx
            )

            cell.font = bold_font
            cell.fill = red_fill

        # --------------------------------------------------------
        # Error Rows
        # --------------------------------------------------------

        for item in failed_rows:

            row_to_write = [
                item["row_no"]
            ]

            for value in item["row_data"]:

                if isinstance(
                    value,
                    datetime
                ):

                    row_to_write.append(
                        value.strftime(
                            "%Y-%m-%d"
                        )
                    )

                else:

                    row_to_write.append(
                        value
                    )

            row_to_write.append(
                item["error_type"]
            )

            row_to_write.append(
                item["error"]
            )

            ws.append(
                row_to_write
            )

        # --------------------------------------------------------
        # Column Width
        # --------------------------------------------------------

        for column in ws.columns:

            max_length = 0

            for cell in column:

                value = str(
                    cell.value or ""
                )

                max_length = max(
                    max_length,
                    len(value)
                )

            ws.column_dimensions[
                column[0].column_letter
            ].width = min(
                max(max_length + 3, 10),
                60
            )

        # --------------------------------------------------------
        # Save Error File
        # --------------------------------------------------------

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"FM_Import_Errors_{timestamp}.xlsx"
        )

        folder_path = get_site_path(
            "private",
            "files"
        )

        os.makedirs(
            folder_path,
            exist_ok=True
        )

        file_path = os.path.join(
            folder_path,
            filename
        )

        wb.save(
            file_path
        )

        # --------------------------------------------------------
        # Create Frappe File
        # --------------------------------------------------------

        file_doc = frappe.new_doc(
            "File"
        )

        file_doc.file_name = filename

        file_doc.is_private = 1

        file_doc.content_type = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

        file_doc.file_url = (
            f"/private/files/{filename}"
        )

        file_doc.attached_to_doctype = (
            self.import_doc.doctype
        )

        file_doc.attached_to_name = (
            self.import_doc.name
        )

        file_doc.insert(
            ignore_permissions=True
        )

        return file_doc.file_url