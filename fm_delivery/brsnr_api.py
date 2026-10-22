"""BRSNR API (fm_delivery) - Frappe / ERPNext 15.

Daily Excel upload -> validation -> snapshot comparison (New / Pending / Cleared)
-> BRSNR Shipment (master) -> BRSNR Shipment History (append-only) -> dashboard.

Business invariants enforced here
---------------------------------
* Comparison key is (Source, ShipmentId). Source is the authorization boundary;
CurrentHub is operational only and never grants access.
* due_month = brsnr_month + 2 calendar months, due_date = last day of due_month.
* LOSS IS FINAL. Once Loss, never Cleared / Pending / Due Soon again, and
loss_date / loss_amount are never lost. loss_amount always comes from
total_price (never gmv).
* History is append-only. Previous snapshot = OPEN rows only (New / Pending).
* Loss reporting is cumulative THROUGH the selected / trend date (loss_date <= date).
"""

import calendar
import hashlib
import hmac
import json
import math
import re
import secrets
import time
from contextlib import contextmanager
from datetime import date, datetime, timedelta

import frappe
from frappe import _
from frappe.utils import (
    cint,
    escape_html,
    flt,
    get_datetime,
    getdate,
    now_datetime,
    strip_html_tags,
    today,
)
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SHIPMENT_DOCTYPE = "BRSNR Shipment"
HISTORY_DOCTYPE = "BRSNR Shipment History"
UPLOAD_DOCTYPE = "BRSNR Daily Upload"
HUB_DOCTYPE = "BRSNR Hub"
ASSIGNMENT_DOCTYPE = "BRSNR Hub Assignment"
INCHARGE_DOCTYPE = "BRSNR Hub Incharge"

REQUIRED_COLUMNS = {"ShipmentId", "Source"}

COLUMN_MAP = {
    "ShipmentId": "shipment_id",
    "Casper ID": "casper_id",
    "Profile ID": "profile_id",
    "Partner": "partner",
    "WM Name": "wm_name",
    "Source": "source",
    "State": "state",
    "Source Type": "source_type",
    "Source Zone": "source_zone",
    "RTO/RVP/FWD": "rtorvpfwd",
    "ForwardReverseType": "forward_reverse_type",
    "Type": "type",
    "ReverseShipmentId": "reverse_shipment_id",
    "TotalPrice": "total_price",
    "Price Range": "price_range",
    "Value Bucket": "value_bucket",
    "gmv": "gmv",
    "DeliveryPincode": "delivery_pincode",
    "DeliveryHub": "delivery_hub",
    "CurrentHub": "current_hub",
    "Status": "status",
    "Final Remarks": "final_remarks",
    "Date From": "date_from",
    "CustomerPromiseDate": "customer_promise_date",
    "LogisticsPromiseDate": "logistics_promise_date",
    "LatestUpdateTime": "latest_update_time",
    "Aging": "aging",
    "Aging Bucket": "aging_bucket",
    "FirstReceiveTime": "first_receive_time",
    "FirstReceiveHub": "first_receive_hub",
    "LastReceiveTime": "last_receive_time",
}
COLUMN_FIELDS = set(COLUMN_MAP.values())
FIELD_TO_COLUMN = {v: k for k, v in COLUMN_MAP.items()}

OPEN_STATUSES = ("New", "Pending")
BRSNR_STATUSES = ("New", "Pending", "Cleared")
CLEARANCE_STATUSES = ("Pending", "Due Soon", "Overdue", "Cleared", "Loss")
RISK_LEVELS = ("Safe", "Watch", "Critical", "Loss")
LIFECYCLE_FIELDS = (
    "brsnr_month",
    "due_month",
    "clearance_status",
    "risk_level",
    "days_remaining",
    "loss_date",
    "loss_amount",
)

DUE_SOON_DAYS = 30
CRITICAL_DAYS = 15

# Auth settings
OTP_TTL_SECONDS = 300
OTP_MAX_FAILED_ATTEMPTS = 5
OTP_LOCKOUT_SECONDS = 900
OTP_RESEND_COOLDOWN_SECONDS = 60
OTP_SENDS_PER_HOUR = 5
OTP_SENDS_PER_IP_PER_HOUR = 30
OTP_VERIFY_PER_IP_PER_15MIN = 40
SESSION_TTL_SECONDS = 8 * 60 * 60

# Dashboard settings
MAX_PAGE_LIMIT = 2000
DEFAULT_TREND_DAYS = 30
MAX_TREND_DAYS = 366
LOSS_REGISTER_LIMIT = 1000
SORTABLE_FIELDS = {
    "aging": "x.aging",
    "total_price": "x.total_price",
    "shipment_id": "x.shipment_id",
    "source": "x.source",
    "current_hub": "x.current_hub",
    "brsnr_status": "x.brsnr_status",
    "clearance_status": "x.clearance_status",
    "days_remaining": "x.days_remaining",
    "loss_amount": "x.loss_amount",
    "due_month": "x.due_month",
}


# ---------------------------------------------------------------------------
# Generic value helpers
# ---------------------------------------------------------------------------

def clean_value(value):
    """Convert Excel values into safe Frappe values (blank -> None)."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def normalize_identifier(value):
    """Normalise ShipmentId / Source (Excel may give 123.0 for numeric ids)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def parse_number(value):
    """Strict number parser. Raises ValueError for anything that is not finite."""
    if isinstance(value, bool):
        raise ValueError("boolean is not a number")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        number = float(str(value).strip().replace(",", ""))
    if math.isnan(number) or math.isinf(number):
        raise ValueError("not a finite number")
    return number


def parse_total_price(value):
    """TotalPrice drives financial LOSS: never silently coerce bad data to 0."""
    if value is None:
        return None
    number = parse_number(value)
    if number < 0:
        raise ValueError("negative TotalPrice")
    return number


_TEXT_FIELDTYPES = {
    "Data", "Small Text", "Text", "Long Text", "Link", "Select", "Dynamic Link",
}


def coerce_value(fieldtype, value):
    """Coerce an Excel cell to the DocField type. Raises ValueError if invalid."""
    if value is None:
        return None

    if fieldtype in ("Float", "Currency", "Percent"):
        return parse_number(value)
    if fieldtype in ("Int", "Long Int"):
        return int(round(parse_number(value)))
    if fieldtype == "Date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        return getdate(str(value))
    if fieldtype == "Datetime":
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day)
        return get_datetime(str(value))
    if fieldtype in _TEXT_FIELDTYPES:
        return normalize_identifier(value)
    return value


def _to_date(value):
    if isinstance(value, datetime):
        return value.date()
    return getdate(value)


def _attr(obj, name):
    if obj is None:
        return None
    if isinstance(obj, dict):
        return obj.get(name)
    return getattr(obj, name, None)


def row_key(row):
    return (str(row.get("source")).strip(), str(row.get("shipment_id")).strip())


def _parse_date(value, label):
    if value in (None, ""):
        return None
    try:
        return getdate(value)
    except Exception:
        frappe.throw(_("Invalid {0}.").format(label))


def _clean_choice(value, allowed, label):
    if value in (None, "", "All", "all"):
        return None
    value = str(value).strip()
    if value not in allowed:
        frappe.throw(_("Invalid {0} filter.").format(label))
    return value


def _normalize_filter_list(value):
    """Accept single values, comma-separated strings, and JSON arrays for multi-select filters."""
    if value is None:
        return []

    if isinstance(value, (list, tuple, set)):
        values = []
        for item in value:
            values.extend(_normalize_filter_list(item))
        return values

    if isinstance(value, str):
        if value.strip() == "":
            return []
        text = value.strip()
        if text.startswith("[") and text.endswith("]"):
            try:
                parsed = json.loads(text)
                if isinstance(parsed, list):
                    return _normalize_filter_list(parsed)
            except Exception:
                pass
        return [part.strip() for part in re.split(r"[,;\n]+", text) if part.strip()]

    return [str(value).strip()] if str(value).strip() else []


def _clean_choices(value, allowed, label):
    """Normalise a multi-select filter to a de-duped list of valid values."""
    values = []
    for item in _normalize_filter_list(value):
        cleaned = _clean_choice(item, allowed, label)
        if label == "clearance status" and cleaned == "Overdue":
            cleaned = "Loss"
        if cleaned and cleaned not in values:
            values.append(cleaned)
    return values


# ---------------------------------------------------------------------------
# Excel reading / validation
# ---------------------------------------------------------------------------

def get_previous_brsnr_date(current_date):
    """Most recent BRSNR date before current_date (uploads skip weekends/holidays)."""
    previous = frappe.db.sql(
        f"SELECT MAX(brsnr_date) FROM `tab{HISTORY_DOCTYPE}` WHERE brsnr_date < %s",
        (current_date,),
    )
    return previous[0][0] if previous and previous[0][0] else None


def get_excel_rows(file_path):
    """Read, type-check and normalise the uploaded XLSX file."""
    field_types = {
        df.fieldname: df.fieldtype for df in frappe.get_meta(SHIPMENT_DOCTYPE).fields
    }

    try:
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception:
        frappe.log_error(title="BRSNR Excel Read Failed", message=frappe.get_traceback())
        frappe.throw(
            _("The uploaded file could not be read. Please upload a valid .xlsx file.")
        )

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)

        try:
            headers = next(rows)
        except StopIteration:
            frappe.throw(_("The Excel file is empty."))

        headers = [str(h).strip() if h is not None else "" for h in headers]

        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            frappe.throw(
                _("Missing required Excel columns: {0}").format(
                    ", ".join(sorted(missing))
                )
            )

        # first occurrence of a column wins; unknown columns are ignored
        index_map = {}
        used_fields = set()
        for index, header in enumerate(headers):
            fieldname = COLUMN_MAP.get(header)
            if fieldname and fieldname in field_types and fieldname not in used_fields:
                index_map[index] = fieldname
                used_fields.add(fieldname)

        result = []
        errors = []

        for row_number, row in enumerate(rows, start=2):
            data = {}
            for index, fieldname in index_map.items():
                if index < len(row):
                    data[fieldname] = clean_value(row[index])

            if all(v is None for v in data.values()):
                continue

            data["shipment_id"] = normalize_identifier(data.get("shipment_id"))
            data["source"] = normalize_identifier(data.get("source"))

            if not data["shipment_id"]:
                errors.append(f"Excel row {row_number}: ShipmentId is missing.")
            elif not data["source"]:
                errors.append(
                    f"Excel row {row_number}: Source is missing for ShipmentId "
                    f"{data['shipment_id']}."
                )
            else:
                for fieldname, value in list(data.items()):
                    if value is None or fieldname in ("shipment_id", "source"):
                        continue
                    try:
                        if fieldname == "total_price":
                            data[fieldname] = parse_total_price(value)
                        else:
                            data[fieldname] = coerce_value(field_types[fieldname], value)
                    except (ValueError, TypeError, OverflowError):
                        errors.append(
                            f"Excel row {row_number}: invalid value in column "
                            f"'{FIELD_TO_COLUMN[fieldname]}' "
                            f"(ShipmentId {data['shipment_id']})."
                        )
                data["_excel_row"] = row_number
                result.append(data)

            if len(errors) >= 200:
                break

        if errors:
            frappe.throw(
                _("Excel validation failed ({0} issue(s) shown up to 50):<br><br>{1}").format(
                    len(errors), "<br>".join(errors[:50])
                )
            )

        return result
    finally:
        workbook.close()


def validate_duplicate_keys(rows):
    """Reject duplicate Source + ShipmentId combinations in the same upload."""
    seen = {}
    duplicates = []

    for row in rows:
        key = row_key(row)
        if key in seen:
            duplicates.append(
                f"{key[0]} + {key[1]} (Excel row {row.get('_excel_row')}, "
                f"first seen at row {seen[key]})"
            )
        else:
            seen[key] = row.get("_excel_row")

    if duplicates:
        frappe.throw(
            _("Duplicate Source + ShipmentId combinations found in Excel:<br><br>{0}").format(
                "<br>".join(duplicates[:50])
            )
        )


def validate_hubs(rows):
    """Source must exactly match an existing BRSNR Hub (name = hub code). One query."""
    valid_hubs = set(frappe.get_all(HUB_DOCTYPE, pluck="name", limit_page_length=0))

    invalid = []
    for row in rows:
        source = str(row.get("source")).strip()
        if source not in valid_hubs:
            invalid.append(
                f"Excel row {row.get('_excel_row')}: Source '{source}' does not exist in BRSNR Hub."
            )

    if invalid:
        frappe.throw(
            _("Invalid Source Hub(s) found ({0}):<br><br>{1}").format(
                len(invalid), "<br>".join(invalid[:100])
            )
        )


def get_previous_snapshot(previous_date):
    """OPEN shipments (New / Pending) of the previous snapshot only.

    Cleared rows must never be carried forward as active records.
    """
    if not previous_date:
        return {}

    rows = frappe.get_all(
        HISTORY_DOCTYPE,
        filters={
            "brsnr_date": previous_date,
            "brsnr_status": ["in", list(OPEN_STATUSES)],
        },
        fields=[
            "name",
            "shipment",
            "shipment_id",
            "source",
            "current_hub",
            "brsnr_status",
            "aging",
            "aging_bucket",
            "final_remarks",
        ],
        limit_page_length=0,
    )

    return {
        (str(r.source).strip(), str(r.shipment_id).strip()): r for r in rows
    }


def get_existing_masters(keys):
    """Set-based master lookup for all keys of this upload (chunked, minimal fields)."""
    wanted = set(keys)
    shipment_ids = sorted({k[1] for k in wanted})
    result = {}

    for start in range(0, len(shipment_ids), 2000):
        chunk = shipment_ids[start:start + 2000]
        rows = frappe.get_all(
            SHIPMENT_DOCTYPE,
            filters={"shipment_id": ["in", chunk]},
            fields=[
                "name", "shipment_id", "source", "status", "current_hub",
                "brsnr_status", "total_price", "brsnr_month", "clearance_status",
                "loss_date", "loss_amount",
            ],
            order_by="creation asc",
            limit_page_length=0,
        )
        for r in rows:
            key = (str(r.source).strip(), str(r.shipment_id).strip())
            if key in wanted:
                result.setdefault(key, r)

    return result


def validate_masters_and_prices(rows, masters, previous_snapshot):
    """TotalPrice must be valid for any shipment that has no stored price, and every
    previously-open shipment must have a master record."""
    errors = []

    for row in rows:
        key = row_key(row)
        master = masters.get(key)

        if key in previous_snapshot and not master:
            errors.append(
                f"Excel row {row.get('_excel_row')}: master record missing for "
                f"{key[0]} + {key[1]} (data inconsistency)."
            )
        elif row.get("total_price") is None and (
            not master or master.total_price is None
        ):
            errors.append(
                f"Excel row {row.get('_excel_row')}: TotalPrice is required for new "
                f"shipment {key[0]} + {key[1]}."
            )

    current_keys = {row_key(r) for r in rows}
    for key in previous_snapshot:
        if key not in current_keys and key not in masters:
            errors.append(
                f"Master record missing for previously open shipment {key[0]} + {key[1]}."
            )

    if errors:
        frappe.throw(
            _("Financial/master validation failed ({0}):<br><br>{1}").format(
                len(errors), "<br>".join(errors[:50])
            )
        )


# ---------------------------------------------------------------------------
# Lifecycle (due date / Loss / risk)
# ---------------------------------------------------------------------------

def get_due_dates(brsnr_month):
    """Return (brsnr_month_first_day, due_month_first_day, due_date)."""
    first = _to_date(brsnr_month).replace(day=1)
    index = first.month - 1 + 2
    year = first.year + index // 12
    month = index % 12 + 1
    due_month = date(year, month, 1)
    due_date = date(year, month, calendar.monthrange(year, month)[1])
    return first, due_month, due_date


def calculate_brsnr_lifecycle(
    brsnr_month,
    is_cleared=False,
    total_price=0,
    as_of_date=None,
    cleared_date=None,
):
    """Pure month-based lifecycle calculation (no stored state considered).

    as_of_date is the BRSNR business date being evaluated (defaults to today).
    Loss amount is TotalPrice, never GMV.
    """
    if not brsnr_month:
        brsnr_month = getdate(today()).replace(day=1)

    first, due_month, due_date = get_due_dates(brsnr_month)
    as_of = _to_date(as_of_date) if as_of_date else getdate(today())
    amount = flt(total_price)

    base = {
        "brsnr_month": first,
        "due_month": due_month,
        "due_date": due_date,
    }

    if is_cleared:
        effective_cleared = _to_date(cleared_date) if cleared_date else as_of
        if effective_cleared > due_date:
            # Cleared late: operationally cleared, financially LOSS.
            return {
                **base,
                "clearance_status": "Loss",
                "risk_level": "Loss",
                "days_remaining": 0,
                "loss_date": due_date,
                "loss_amount": amount,
            }
        return {
            **base,
            "clearance_status": "Cleared",
            "risk_level": "Safe",
            "days_remaining": max(0, (due_date - effective_cleared).days),
            "loss_date": None,
            "loss_amount": 0,
        }

    if as_of > due_date:
        return {
            **base,
            "clearance_status": "Loss",
            "risk_level": "Loss",
            "days_remaining": 0,
            "loss_date": due_date,
            "loss_amount": amount,
        }

    days_remaining = (due_date - as_of).days

    if days_remaining <= CRITICAL_DAYS:
        status, risk = "Due Soon", "Critical"
    elif days_remaining <= DUE_SOON_DAYS:
        status, risk = "Due Soon", "Watch"
    else:
        status, risk = "Pending", "Safe"

    return {
        **base,
        "clearance_status": status,
        "risk_level": risk,
        "days_remaining": days_remaining,
        "loss_date": None,
        "loss_amount": 0,
    }


def resolve_lifecycle(
    existing,
    brsnr_month,
    total_price,
    as_of_date,
    is_cleared=False,
    cleared_date=None,
):
    """Lifecycle with the LOSS-IS-FINAL rule applied on top of the stored state.

    `existing` is the current master row (dict/doc) or None.
    """
    price_missing = total_price is None
    calc = calculate_brsnr_lifecycle(
        brsnr_month,
        is_cleared=is_cleared,
        total_price=total_price or 0,
        as_of_date=as_of_date,
        cleared_date=cleared_date,
    )

    existing_loss_date = _attr(existing, "loss_date")
    already_loss = _attr(existing, "clearance_status") == "Loss" or bool(existing_loss_date)

    if not already_loss and calc["clearance_status"] != "Loss":
        calc["price_missing"] = False
        return calc

    if already_loss:
        loss_date = _to_date(existing_loss_date) if existing_loss_date else calc["due_date"]
        stored_loss_amount = _attr(existing, "loss_amount")
        loss_amount = (
            flt(stored_loss_amount)
            if stored_loss_amount is not None
            else flt(total_price)
        )
    else:
        loss_date = calc["loss_date"]
        loss_amount = flt(total_price)

    return {
        "brsnr_month": calc["brsnr_month"],
        "due_month": calc["due_month"],
        "due_date": calc["due_date"],
        "clearance_status": "Loss",
        "risk_level": "Loss",
        "days_remaining": 0,
        "loss_date": loss_date,
        "loss_amount": loss_amount,
        "price_missing": price_missing and not loss_amount,
    }


def apply_brsnr_lifecycle(shipment, is_cleared=False, as_of_date=None, cleared_date=None):
    """Apply lifecycle fields to a BRSNR Shipment document. LOSS is permanent."""
    brsnr_month = getattr(shipment, "brsnr_month", None) or getdate(
        getattr(shipment, "brsnr_date", None) or today()
    ).replace(day=1)

    lifecycle = resolve_lifecycle(
        shipment,
        brsnr_month,
        getattr(shipment, "total_price", None),
        as_of_date or getattr(shipment, "brsnr_date", None) or today(),
        is_cleared=is_cleared,
        cleared_date=cleared_date,
    )

    for fieldname in LIFECYCLE_FIELDS:
        setattr(shipment, fieldname, lifecycle[fieldname])


def _guard_loss_final(master, values):
    """Defensive invariant: an existing Loss can never be altered or reversed.

    The lifecycle logic already guarantees this; the guard aborts (and rolls back)
    the whole import if a future code change ever breaks it.
    """
    if not master:
        return

    old_loss_date = _attr(master, "loss_date")
    if _attr(master, "clearance_status") != "Loss" and not old_loss_date:
        return

    old_amount = _attr(master, "loss_amount")
    new_date = values.get("loss_date")
    broken = (
        values.get("clearance_status") != "Loss"
        or not new_date
        or (old_loss_date and getdate(new_date) != getdate(old_loss_date))
        or (
            old_amount is not None
            and flt(values.get("loss_amount")) != flt(old_amount)
        )
    )
    if broken:
        frappe.throw(
            _("Loss reversal blocked for {0} + {1}.").format(
                _attr(master, "source"), _attr(master, "shipment_id")
            )
        )


# ---------------------------------------------------------------------------
# Record builders (set-based importer)
# ---------------------------------------------------------------------------

def build_history_record(
    shipment_name,
    row,
    brsnr_date,
    upload_name,
    brsnr_status,
    previous_status=None,
    previous_hub=None,
):
    """Append-only snapshot row for BRSNR Shipment History (a dict, inserted in bulk)."""
    return {
        "shipment": shipment_name,
        "shipment_id": row.get("shipment_id"),
        "brsnr_date": brsnr_date,
        "daily_upload": upload_name,
        "source": row.get("source"),
        "current_hub": row.get("current_hub"),
        "status": row.get("status"),
        "brsnr_status": brsnr_status,
        "aging": row.get("aging"),
        "aging_bucket": row.get("aging_bucket"),
        "final_remarks": row.get("final_remarks"),
        "previous_hub": previous_hub,
        "previous_status": previous_status,
    }


def build_master_values(
    row,
    master,
    brsnr_date,
    upload_name,
    brsnr_status,
    previous_status=None,
    previous_hub=None,
    cleared_date=None,
):
    """Values for a New/Pending master (create or update).

    Blank Excel values never overwrite existing master data.
    """
    values = {
        f: v for f, v in row.items() if f in COLUMN_FIELDS and v is not None
    }
    values.update(
        {
            "source": row["source"],
            "shipment_id": row["shipment_id"],
            "brsnr_status": brsnr_status,
            "brsnr_date": brsnr_date,
            "daily_upload": upload_name,
            "previous_status": previous_status,
            "previous_hub": previous_hub,
            "cleared_date": cleared_date,
        }
    )

    price = row.get("total_price")
    if price is None:
        price = _attr(master, "total_price")

    brsnr_month = _attr(master, "brsnr_month") or brsnr_date.replace(day=1)

    lifecycle = resolve_lifecycle(
        master,
        brsnr_month,
        price,
        brsnr_date,
        is_cleared=(brsnr_status == "Cleared"),
        cleared_date=cleared_date,
    )
    for fieldname in LIFECYCLE_FIELDS:
        values[fieldname] = lifecycle[fieldname]

    return values, lifecycle


def build_cleared_master_values(master, previous, brsnr_date, upload_name):
    """Values for a shipment that disappeared from today's file (Cleared event).

    Only lifecycle/status fields are touched; operational data is preserved.
    """
    brsnr_month = _attr(master, "brsnr_month") or brsnr_date.replace(day=1)

    lifecycle = resolve_lifecycle(
        master,
        brsnr_month,
        _attr(master, "total_price"),
        brsnr_date,
        is_cleared=True,
        cleared_date=brsnr_date,
    )

    values = {
        "brsnr_status": "Cleared",
        "brsnr_date": brsnr_date,
        "daily_upload": upload_name,
        "previous_status": previous.brsnr_status,
        "previous_hub": previous.current_hub,
        "cleared_date": brsnr_date,
    }
    for fieldname in LIFECYCLE_FIELDS:
        values[fieldname] = lifecycle[fieldname]

    return values, lifecycle


def _insert_chunks(doctype, columns, values, chunk_size=500):
    """Parameterised multi-row INSERT (never string-formats cell data into SQL)."""
    placeholder = "(" + ", ".join(["%s"] * len(columns)) + ")"
    column_sql = ", ".join(f"`{c}`" for c in columns)

    for start in range(0, len(values), chunk_size):
        part = values[start:start + chunk_size]
        query = (
            f"INSERT INTO `tab{doctype}` ({column_sql}) VALUES "
            + ", ".join([placeholder] * len(part))
        )
        frappe.db.sql(query, [v for row in part for v in row])


def insert_records(doctype, records):
    """Insert many records efficiently; returns the list of document names.

    Hash-named doctypes use a parameterised bulk INSERT. Doctypes with any other
    naming rule fall back to normal document inserts so naming stays correct.
    """
    if not records:
        return []

    meta = frappe.get_meta(doctype)
    clean = [
        {k: v for k, v in rec.items() if v is not None and meta.has_field(k)}
        for rec in records
    ]

    autoname = (meta.autoname or "").strip().lower()
    names = [None] * len(clean)

    if autoname in ("", "hash"):
        now = now_datetime()
        user = frappe.session.user
        groups = {}

        for index, rec in enumerate(clean):
            names[index] = frappe.generate_hash(length=12)
            groups.setdefault(tuple(sorted(rec)), []).append(index)

        for signature, indexes in groups.items():
            columns = [
                "name", "creation", "modified", "modified_by", "owner",
                "docstatus", "idx", *signature,
            ]
            values = [
                [names[i], now, now, user, user, 0, 0, *[clean[i][f] for f in signature]]
                for i in indexes
            ]
            _insert_chunks(doctype, columns, values)
    else:
        for index, rec in enumerate(clean):
            doc = frappe.new_doc(doctype)
            doc.update(rec)
            doc.insert(ignore_permissions=True)
            names[index] = doc.name

    return names


# ---------------------------------------------------------------------------
# Import engine
# ---------------------------------------------------------------------------

@contextmanager
def brsnr_import_lock():
    """Global, cross-process import lock (snapshot comparison is order dependent)."""
    if frappe.db.db_type == "mariadb":
        lock_name = "brsnr_import_" + hashlib.sha1(
            str(frappe.local.site).encode()
        ).hexdigest()[:24]

        got = frappe.db.sql("SELECT GET_LOCK(%s, 0)", (lock_name,))
        if not got or got[0][0] != 1:
            frappe.throw(
                _("Another BRSNR import is currently running. Please try again after it finishes.")
            )
        try:
            yield
        finally:
            try:
                frappe.db.sql("SELECT RELEASE_LOCK(%s)", (lock_name,))
            except Exception:
                pass
    else:
        from frappe.utils.synchronization import filelock

        with filelock("brsnr_import", timeout=1):
            yield


def _resolve_file_path(upload):
    file_name = frappe.db.get_value("File", {"file_url": upload.import_file}, "name")
    if not file_name:
        frappe.throw(_("The attached Excel file could not be found."))

    file_doc = frappe.get_doc("File", file_name)
    if not (file_doc.file_name or "").lower().endswith(".xlsx"):
        frappe.throw(_("Only .xlsx files are supported."))

    return file_doc.get_full_path()


def _assert_order_and_idempotency(upload, current_date):
    """Prevent duplicate processing and out-of-order snapshots."""
    if frappe.db.exists(HISTORY_DOCTYPE, {"daily_upload": upload.name}):
        frappe.throw(
            _("History already exists for this upload. Refusing to create duplicate history.")
        )

    other = frappe.db.sql(
        f"""
        SELECT name FROM `tab{UPLOAD_DOCTYPE}`
        WHERE status = 'Completed' AND name != %s AND brsnr_date >= %s
        LIMIT 1
        """,
        (upload.name, current_date),
    )
    if other:
        frappe.throw(
            _("BRSNR date {0} (or a later date) has already been processed by upload {1}.").format(
                current_date, other[0][0]
            )
        )

    latest = frappe.db.sql(f"SELECT MAX(brsnr_date) FROM `tab{HISTORY_DOCTYPE}`")
    latest = latest[0][0] if latest and latest[0][0] else None
    if latest and getdate(latest) >= current_date:
        if getdate(latest) == current_date:
            frappe.throw(_("BRSNR data for {0} has already been processed.").format(current_date))
        frappe.throw(
            _(
                "Out-of-order upload: data up to {0} is already processed. "
                "Uploads must be processed in chronological order."
            ).format(latest)
        )


def _run_import(upload, current_date):
    _assert_order_and_idempotency(upload, current_date)

    rows = get_excel_rows(_resolve_file_path(upload))
    if not rows:
        frappe.throw(_("No valid records found in the Excel file."))

    validate_duplicate_keys(rows)
    validate_hubs(rows)

    previous_date = get_previous_brsnr_date(current_date)
    previous_snapshot = get_previous_snapshot(previous_date)

    current_keys = {row_key(r) for r in rows}
    masters = get_existing_masters(current_keys | set(previous_snapshot))

    validate_masters_and_prices(rows, masters, previous_snapshot)

    meta = frappe.get_meta(SHIPMENT_DOCTYPE)
    new_entries = []        # (master values, history record)
    updates = []            # (master name, values)
    history_records = []    # history rows for already-existing masters
    price_warnings = []

    new_count = pending_count = cleared_count = 0

    # --- today's records: New / Pending ------------------------------------
    for row in rows:
        key = row_key(row)
        previous = previous_snapshot.get(key)
        master = masters.get(key)

        if previous:
            brsnr_status = "Pending"
            previous_status = previous.brsnr_status
            previous_hub = previous.current_hub
            pending_count += 1
        else:
            brsnr_status = "New"
            # A shipment that reappears after being Cleared keeps its master and
            # its financial lifecycle (Loss stays Loss); it is flagged via
            # previous_status for traceability.
            reappeared = bool(master and master.brsnr_status == "Cleared")
            previous_status = "Cleared" if reappeared else None
            previous_hub = master.current_hub if reappeared else None
            new_count += 1

        values, lifecycle = build_master_values(
            row, master, current_date, upload.name, brsnr_status,
            previous_status, previous_hub, cleared_date=None,
        )
        if lifecycle.get("price_missing"):
            price_warnings.append(f"{key[0]}+{key[1]}")

        history = build_history_record(
            master.name if master else None,
            row, current_date, upload.name, brsnr_status,
            previous_status, previous_hub,
        )

        if master:
            _guard_loss_final(master, values)
            updates.append((master.name, values))
            history_records.append(history)
        else:
            new_entries.append((values, history))

    # --- previously open, missing today: Cleared ---------------------------
    for key, previous in previous_snapshot.items():
        if key in current_keys:
            continue

        master = masters[key]
        values, lifecycle = build_cleared_master_values(
            master, previous, current_date, upload.name
        )
        _guard_loss_final(master, values)
        if lifecycle.get("price_missing"):
            price_warnings.append(f"{key[0]}+{key[1]}")

        updates.append((master.name, values))

        snapshot_row = {
            "shipment_id": previous.shipment_id,
            "source": previous.source,
            "current_hub": previous.current_hub,
            "status": master.status,
            "aging": previous.aging,
            "aging_bucket": previous.aging_bucket,
            "final_remarks": previous.final_remarks,
        }
        history_records.append(
            build_history_record(
                master.name, snapshot_row, current_date, upload.name, "Cleared",
                previous.brsnr_status, previous.current_hub,
            )
        )
        cleared_count += 1

    # --- writes (single transaction, committed by the caller) --------------
    if new_entries:
        names = insert_records(SHIPMENT_DOCTYPE, [e[0] for e in new_entries])
        for (values, history), name in zip(new_entries, names):
            history["shipment"] = name
            history_records.append(history)

    for name, values in updates:
        frappe.db.set_value(
            SHIPMENT_DOCTYPE,
            name,
            {k: v for k, v in values.items() if meta.has_field(k)},
        )

    insert_records(HISTORY_DOCTYPE, history_records)

    frappe.db.set_value(
        UPLOAD_DOCTYPE,
        upload.name,
        {
            "total_records": len(rows),
            "new_records": new_count,
            "pending_records": pending_count,
            "cleared_records": cleared_count,
            "status": "Completed",
            "error_message": None,
        },
    )

    if price_warnings:
        frappe.log_error(
            title="BRSNR Loss with missing TotalPrice",
            message="Loss recorded with no TotalPrice for: " + ", ".join(price_warnings[:200]),
        )

    return {
        "success": True,
        "upload": upload.name,
        "brsnr_date": str(current_date),
        "previous_date": str(previous_date) if previous_date else None,
        "total_records": len(rows),
        "new_records": new_count,
        "pending_records": pending_count,
        "cleared_records": cleared_count,
    }


def _mark_failed(upload_name, exc):
    if isinstance(exc, frappe.ValidationError):
        message = strip_html_tags(str(exc).replace("<br>", "\n"))
        reference = None
    else:
        log = frappe.log_error(
            title="BRSNR Import Failed",
            message=frappe.get_traceback(),
            reference_doctype=UPLOAD_DOCTYPE,
            reference_name=upload_name,
        )
        reference = getattr(log, "name", None)
        message = "Unexpected error while importing. See Error Log" + (
            f" {reference}." if reference else "."
        )

    frappe.db.set_value(
        UPLOAD_DOCTYPE, upload_name, {"status": "Failed", "error_message": message[:5000]}
    )
    frappe.db.commit()
    return reference


@frappe.whitelist()
def process_brsnr_upload(upload_name):
    """Main BRSNR import engine.

    Excel -> validate -> compare with previous OPEN snapshot -> New / Pending /
    Cleared -> BRSNR Shipment (master) -> BRSNR Shipment History (append-only).

    All data writes happen in ONE transaction that is committed together with the
    Completed status; any error rolls everything back and marks the upload Failed.
    """
    if not upload_name:
        frappe.throw(_("Upload is required."))

    frappe.has_permission(UPLOAD_DOCTYPE, "write", doc=upload_name, throw=True)

    with brsnr_import_lock():
        frappe.db.commit()  # fresh snapshot after acquiring the lock

        upload = frappe.get_doc(UPLOAD_DOCTYPE, upload_name)

        if upload.status == "Completed":
            return {
                "success": True,
                "already_processed": True,
                "upload": upload.name,
                "brsnr_date": str(getdate(upload.brsnr_date)),
                "total_records": upload.total_records,
                "new_records": upload.new_records,
                "pending_records": upload.pending_records,
                "cleared_records": upload.cleared_records,
            }

        if not upload.import_file:
            frappe.throw(_("Please attach the BRSNR Excel file."))
        if not upload.brsnr_date:
            frappe.throw(_("Please enter BRSNR Date."))

        current_date = getdate(upload.brsnr_date)
        if current_date > getdate(today()):
            frappe.throw(_("BRSNR Date cannot be in the future."))

        # A 'Processing' status here can only be a crashed run: the global lock is
        # held by us, and data is only committed together with 'Completed'.
        frappe.db.set_value(
            UPLOAD_DOCTYPE,
            upload.name,
            {
                "status": "Processing",
                "uploaded_by": frappe.session.user,
                "upload_time": now_datetime(),
                "error_message": None,
            },
        )
        frappe.db.commit()

        try:
            result = _run_import(upload, current_date)
            frappe.db.commit()
            return result
        except Exception as exc:
            frappe.db.rollback()
            reference = _mark_failed(upload.name, exc)
            if isinstance(exc, frappe.ValidationError):
                raise
            frappe.throw(
                _("The BRSNR import failed unexpectedly. Reference: {0}").format(reference or "-")
            )


# ---------------------------------------------------------------------------
# Portal identity
# ---------------------------------------------------------------------------

def _normalize_email(email):
    return (str(email or "")).strip().lower()[:140]


def get_portal_identity(email):
    """Resolve an allowed BRSNR portal identity.

    Director: enabled Frappe User holding the BRSNR Director role (checked first,
    as it is the superset). Hub Incharge: active BRSNR Hub Incharge by email.
    """
    email = _normalize_email(email)
    if not email:
        return None

    user = frappe.db.get_value(
        "User",
        {"name": email, "enabled": 1},
        ["name", "full_name", "email"],
        as_dict=True,
    )
    if user and "BRSNR Director" in frappe.get_roles(user.name):
        return {
            "type": "Director",
            "name": user.name,
            "display_name": user.full_name or user.email or user.name,
            "email": (user.email or user.name).lower(),
        }

    incharge = frappe.db.get_value(
        INCHARGE_DOCTYPE,
        {"email": email, "active": 1},
        ["name", "incharge_name", "email"],
        as_dict=True,
    )
    if incharge:
        return {
            "type": "Incharge",
            "name": incharge.name,
            "display_name": incharge.incharge_name,
            "email": (incharge.email or email).lower(),
        }

    return None


def get_authorized_hubs(identity, as_of=None):
    """Server-side authorization scope. Returns (hub_names, assignments, hub_options).

    Director -> ALL hubs (including inactive).
    Incharge -> hubs with an assignment that is active and valid on `as_of`.
    `Source` is the authorization field; CurrentHub is never used.
    """
    as_of = getdate(as_of or today())

    if identity["type"] == "Director":
        hubs = frappe.get_all(
            HUB_DOCTYPE,
            fields=["name", "hub_name", "hub_code", "active"],
            order_by="hub_name asc",
            limit_page_length=0,
        )
        return [h.name for h in hubs], [], [dict(h) for h in hubs]

    assignments = frappe.db.sql(
        f"""
        SELECT hub, from_date, to_date
        FROM `tab{ASSIGNMENT_DOCTYPE}`
        WHERE hub_incharge = %(incharge)s
        AND active = 1
        AND (from_date IS NULL OR from_date <= %(d)s)
        AND (to_date IS NULL OR to_date >= %(d)s)
        """,
        {"incharge": identity["name"], "d": as_of},
        as_dict=True,
    )

    hub_names = []
    for a in assignments:
        if a.hub and a.hub not in hub_names:
            hub_names.append(a.hub)

    options = []
    if hub_names:
        options = frappe.get_all(
            HUB_DOCTYPE,
            filters={"name": ["in", hub_names]},
            fields=["name", "hub_name", "hub_code", "active"],
            order_by="hub_name asc",
            limit_page_length=0,
        )
        options = [dict(h) for h in options]

    return hub_names, [dict(a) for a in assignments], options


# ---------------------------------------------------------------------------
# OTP + session tokens
# ---------------------------------------------------------------------------

def _auth_secret():
    try:
        from frappe.utils.password import get_encryption_key

        return str(get_encryption_key()).encode()
    except Exception:
        return (str(frappe.local.site) + ":brsnr").encode()


def _email_key(email):
    return hashlib.sha256(email.encode()).hexdigest()[:32]


def _otp_hash(email, otp):
    return hmac.new(_auth_secret(), f"{email}:{otp}".encode(), hashlib.sha256).hexdigest()


def _token_hash(token):
    return hashlib.sha256(token.encode()).hexdigest()


def _hit(key, window_seconds):
    """Atomic counter with a fixed window (same primitive Frappe's rate limiter uses)."""
    cache = frappe.cache()
    redis_key = cache.make_key(key)
    count = cache.incrby(redis_key, 1)
    if count == 1 or cache.ttl(redis_key) < 0:
        cache.expire(redis_key, window_seconds)
    return int(count)


def _peek(key):
    cache = frappe.cache()
    value = cache.get(cache.make_key(key))
    if isinstance(value, bytes):
        value = value.decode()
    return cint(value)


def _request_ip():
    return getattr(frappe.local, "request_ip", None) or "unknown"


@frappe.whitelist(allow_guest=True, methods=["POST"])
def send_login_otp(email=None):
    """Send an OTP to an active BRSNR Incharge or Director.

    The response is identical for eligible and ineligible addresses.
    """
    generic = {
        "success": True,
        "message": _(
            "If the account is eligible, a verification code has been sent to the registered email."
        ),
    }

    email = _normalize_email(email)
    if not email:
        return generic

    if _hit(f"brsnr_otp_ip:{_request_ip()}", 3600) > OTP_SENDS_PER_IP_PER_HOUR:
        frappe.throw(_("Too many requests. Please try again later."))

    key_id = _email_key(email)

    if _hit(f"brsnr_otp_cd:{key_id}", OTP_RESEND_COOLDOWN_SECONDS) > 1:
        frappe.throw(_("Please wait a minute before requesting another code."))

    if _hit(f"brsnr_otp_hr:{key_id}", 3600) > OTP_SENDS_PER_HOUR:
        frappe.throw(_("Too many requests. Please try again later."))

    identity = get_portal_identity(email)
    if not identity:
        return generic

    otp = f"{secrets.randbelow(10 ** 6):06d}"
    frappe.cache().set_value(
        f"brsnr_otp:{key_id}",
        {"hash": _otp_hash(email, otp), "expires_at": time.time() + OTP_TTL_SECONDS},
        expires_in_sec=OTP_TTL_SECONDS,
    )

    try:
        frappe.sendmail(
            recipients=[identity["email"]],
            subject="BRSNR Secure Login — One-Time Verification Code",
            message=f"""
            <div style="font-family:Arial,sans-serif;max-width:620px;margin:30px auto;padding:35px;background:#fff;border:1px solid #e8edf3;border-radius:18px;color:#172033">
                <div style="font-size:28px;font-weight:800;color:#173f72">BRSNR</div>
                <div style="font-size:11px;color:#7a8699;letter-spacing:1.5px;margin-top:4px">SECURE OPERATIONS PLATFORM</div>
                <hr style="border:0;border-top:1px solid #edf0f5;margin:25px 0">
                <p>Hello <strong>{escape_html(identity['display_name'] or '')}</strong>,</p>
                <p>Use the following verification code to access the BRSNR Dashboard:</p>
                <div style="text-align:center;background:#f7faff;border:1px solid #dce8f7;border-radius:15px;padding:25px;margin:25px 0">
                    <div style="font-size:11px;color:#738096;letter-spacing:2px">ONE-TIME PASSWORD</div>
                    <div style="font-size:40px;font-weight:800;letter-spacing:9px;color:#173f72;margin:15px 0 8px">{otp}</div>
                    <div style="font-size:12px;color:#7b8798">Valid for <strong>5 minutes</strong>.</div>
                </div>
                <p style="font-size:12px;color:#667085">Never share this OTP with anyone.</p>
                <p style="font-size:11px;color:#98a2b3;margin-top:30px">Wercatalyst Ventures Pvt Ltd. · Automated Security Email</p>
            </div>
            """,
            now=True,
        )
    except Exception:
        # never include the OTP / message body in logs
        frappe.log_error(title="BRSNR OTP email failed", message=frappe.get_traceback())

    return generic


def _issue_session(identity):
    token = secrets.token_urlsafe(32)
    now = time.time()
    frappe.cache().set_value(
        f"brsnr_session:{_token_hash(token)}",
        {
            "email": identity["email"].lower(),
            "role": identity["type"],
            "issued_at": now,
            "expires_at": now + SESSION_TTL_SECONDS,
        },
        expires_in_sec=SESSION_TTL_SECONDS,
    )
    return token


@frappe.whitelist(allow_guest=True, methods=["POST"])
def verify_login_otp(email=None, otp=None):
    """Verify the OTP and issue a random, server-side-hashed, expiring session token."""
    email = _normalize_email(email)
    otp = str(otp or "").strip()

    if _hit(f"brsnr_verify_ip:{_request_ip()}", OTP_LOCKOUT_SECONDS) > OTP_VERIFY_PER_IP_PER_15MIN:
        frappe.throw(_("Too many attempts. Please try again later."))

    key_id = _email_key(email)
    fail_key = f"brsnr_otp_fail:{key_id}"
    otp_key = f"brsnr_otp:{key_id}"

    if _peek(fail_key) >= OTP_MAX_FAILED_ATTEMPTS:
        frappe.throw(_("Too many failed attempts. Please try again later."))

    identity = get_portal_identity(email)
    record = frappe.cache().get_value(otp_key)

    valid = False
    if (
        identity
        and isinstance(record, dict)
        and record.get("expires_at", 0) > time.time()
        and re.fullmatch(r"\d{6}", otp)
    ):
        valid = hmac.compare_digest(str(record.get("hash", "")), _otp_hash(email, otp))

    if not valid:
        if _hit(fail_key, OTP_LOCKOUT_SECONDS) >= OTP_MAX_FAILED_ATTEMPTS:
            frappe.cache().delete_value(otp_key)
        frappe.throw(_("Invalid or expired verification code."))

    # single use
    frappe.cache().delete_value(otp_key)
    frappe.cache().delete_value(fail_key)

    token = _issue_session(identity)
    _hubs, assignments, _options = get_authorized_hubs(identity, today())

    return {
        "success": True,
        "role": identity["type"],
        "incharge": identity["name"],
        "incharge_name": identity["display_name"],
        "email": identity["email"],
        "assignments": assignments,
        "token": token,
        "token_type": "Bearer",
        "expires_in": SESSION_TTL_SECONDS,
    }


def _extract_token(token=None):
    """Token sources: explicit param, X-BRSNR-Token header, Authorization: Bearer."""
    if token:
        return str(token).strip()

    header = frappe.get_request_header("X-BRSNR-Token")
    if header:
        return header.strip()

    auth = frappe.get_request_header("Authorization") or ""
    parts = auth.split(None, 1)
    if len(parts) == 2 and parts[0].lower() == "bearer":
        return parts[1].strip()

    return None


def require_portal_identity(token=None):
    """Validate the session token on EVERY protected call and re-check the account."""
    raw = _extract_token(token)
    unauthorized = _("Your session has expired. Please sign in again.")

    if not raw or len(raw) > 200:
        frappe.throw(unauthorized, frappe.AuthenticationError)

    cache_key = f"brsnr_session:{_token_hash(raw)}"
    record = frappe.cache().get_value(cache_key)

    if not isinstance(record, dict) or record.get("expires_at", 0) < time.time():
        frappe.cache().delete_value(cache_key)
        frappe.throw(unauthorized, frappe.AuthenticationError)

    identity = get_portal_identity(record.get("email"))
    if not identity or identity["type"] != record.get("role"):
        frappe.cache().delete_value(cache_key)
        frappe.throw(unauthorized, frappe.AuthenticationError)

    return identity


@frappe.whitelist(allow_guest=True, methods=["POST"])
def logout_portal(token=None):
    """Revoke the current session token."""
    raw = _extract_token(token)
    if raw:
        frappe.cache().delete_value(f"brsnr_session:{_token_hash(raw)}")
    return {"success": True}


@frappe.whitelist(allow_guest=True)
def get_portal_session(token=None):
    """Validate a stored token and return the caller's identity/scope (page reload)."""
    identity = require_portal_identity(token)
    _hubs, assignments, options = get_authorized_hubs(identity, today())
    return {
        "success": True,
        "role": identity["type"],
        "incharge": identity["name"],
        "incharge_name": identity["display_name"],
        "email": identity["email"],
        "assignments": assignments,
        "hub_options": options,
    }


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

_BM = (
    "DATE_SUB(COALESCE(s.brsnr_month, h.brsnr_date), "
    "INTERVAL (DAYOFMONTH(COALESCE(s.brsnr_month, h.brsnr_date)) - 1) DAY)"
)

_LOSS_COND = "b.brsnr_date > COALESCE(b.master_loss_date, b.due_date)"


def _snapshot_sql(date_cond):
    """Point-in-time snapshot built from HISTORY (+ master for static/financial data).

    Clearance status / risk / days_remaining are computed AS OF each history row's
    own brsnr_date, so historical dates are not faked from the current master.
    date_cond is a trusted constant (never user input).
    """
    return f"""
    (SELECT b.*,
        CASE
            WHEN {_LOSS_COND} THEN 'Loss'
            WHEN b.brsnr_status = 'Cleared' THEN 'Cleared'
            WHEN DATEDIFF(b.due_date, b.brsnr_date) <= {DUE_SOON_DAYS} THEN 'Due Soon'
            ELSE 'Pending'
        END AS clearance_status,
        CASE
            WHEN {_LOSS_COND} THEN 'Loss'
            WHEN b.brsnr_status = 'Cleared' THEN 'Safe'
            WHEN DATEDIFF(b.due_date, b.brsnr_date) <= {CRITICAL_DAYS} THEN 'Critical'
            WHEN DATEDIFF(b.due_date, b.brsnr_date) <= {DUE_SOON_DAYS} THEN 'Watch'
            ELSE 'Safe'
        END AS risk_level,
        CASE
            WHEN {_LOSS_COND} THEN 0
            ELSE GREATEST(0, DATEDIFF(b.due_date, b.brsnr_date))
        END AS days_remaining,
        CASE
            WHEN {_LOSS_COND}
            THEN COALESCE(b.master_loss_amount, b.total_price, 0)
            ELSE 0
        END AS loss_amount
    FROM (
        SELECT
            h.name AS history_name, h.shipment, h.shipment_id, h.source,
            h.current_hub, h.status, h.brsnr_status, h.aging, h.aging_bucket,
            h.final_remarks, h.previous_hub, h.previous_status, h.brsnr_date,
            s.total_price,
            s.loss_date AS master_loss_date,
            s.loss_amount AS master_loss_amount,
            {_BM} AS brsnr_month,
            DATE_ADD({_BM}, INTERVAL 2 MONTH) AS due_month,
            LAST_DAY(DATE_ADD({_BM}, INTERVAL 2 MONTH)) AS due_date
        FROM (
            SELECT h.*,
                ROW_NUMBER() OVER (
                    PARTITION BY h.source, h.shipment_id, h.brsnr_date
                    ORDER BY h.creation DESC, h.name DESC
                ) AS history_rank
            FROM `tab{HISTORY_DOCTYPE}` h
            WHERE {date_cond} AND h.source IN %(hubs)s
        ) h
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = h.shipment
        WHERE h.history_rank = 1
    ) b
    ) x
    """


def _loss_register_sql():
    """Cumulative final Loss rows from masters, with historical snapshot details."""
    loss_month = (
        "DATE_SUB(COALESCE(m.brsnr_month, h.brsnr_date, m.brsnr_date), "
        "INTERVAL (DAYOFMONTH(COALESCE(m.brsnr_month, h.brsnr_date, m.brsnr_date)) - 1) DAY)"
    )
    due_month = f"DATE_ADD({loss_month}, INTERVAL 2 MONTH)"
    due_date = f"LAST_DAY({due_month})"
    return f"""
    (
        SELECT
            m.name, m.name AS shipment, m.shipment_id, m.source,
            COALESCE(h.current_hub, m.current_hub) AS current_hub,
            COALESCE(h.brsnr_status, m.brsnr_status) AS brsnr_status,
            COALESCE(h.brsnr_date, m.brsnr_date) AS brsnr_date,
            m.total_price,
            COALESCE(m.loss_amount, m.total_price, 0) AS loss_amount,
            m.partner, m.casper_id,
            COALESCE(m.loss_date, {due_date}) AS loss_date,
            {loss_month} AS brsnr_month,
            {due_month} AS due_month
        FROM `tab{SHIPMENT_DOCTYPE}` m
        LEFT JOIN (
            SELECT ranked.*
            FROM (
                SELECT h.*,
                    ROW_NUMBER() OVER (
                        PARTITION BY h.source, h.shipment_id
                        ORDER BY h.brsnr_date DESC, h.creation DESC, h.name DESC
                    ) AS history_rank
                FROM `tab{HISTORY_DOCTYPE}` h
                WHERE h.brsnr_date <= %(d)s AND h.source IN %(hubs)s
            ) ranked
            WHERE ranked.history_rank = 1
        ) h ON h.shipment = m.name
        WHERE m.source IN %(hubs)s
        AND (
            (m.loss_date IS NOT NULL AND m.loss_date <= %(d)s)
            OR (
                m.clearance_status = 'Loss'
                AND m.loss_date IS NULL
                AND %(d)s > {due_date}
            )
        )
    ) l
    """


_SNAP_ONE = _snapshot_sql("h.brsnr_date = %(d)s")
_SNAP_TREND = _snapshot_sql("h.brsnr_date BETWEEN %(tf)s AND %(d)s")


def _sql(query, params):
    return frappe.db.sql(query, params, as_dict=True)


def _empty_summary():
    return {
        "total": 0, "new": 0, "pending": 0, "cleared": 0,
        "clearance_percentage": 0, "due_soon": 0, "loss": 0,
        "loss_value": 0, "avg_aging": 0, "average_aging": 0,
        "watch": 0, "watch_value": 0, "critical": 0, "critical_value": 0,
    }


def _empty_dashboard(identity, hubs, snapshot_date, latest_date, hub_options=None):
    empty_risk = {
        level: {"count": 0, "amount": 0}
        for level in ("Watch", "Critical", "Loss")
    }
    empty_scorecard_totals = {
        "total": 0, "new_count": 0, "pending_count": 0, "cleared_count": 0,
        "due_soon_count": 0, "watch_count": 0, "critical_count": 0,
        "new_amount": 0, "pending_amount": 0, "cleared_amount": 0,
        "due_soon_amount": 0, "watch_amount": 0, "critical_amount": 0,
        "total_amount": 0, "loss_count": 0, "loss_amount": 0,
        "loss_value": 0, "avg_aging": 0,
    }
    return {
        "success": True,
        "role": identity["type"],
        "incharge": identity["display_name"],
        "email": identity["email"],
        "hubs": hubs,
        "hub_options": hub_options or [],
        "brsnr_date": snapshot_date,
        "latest_brsnr_date": latest_date,
        "summary": _empty_summary(),
        "risk_summary": empty_risk,
        "records": [],
        "records_total": 0,
        "trend": [],
        "hub_scorecard": [],
        "hub_scorecard_totals": empty_scorecard_totals,
        "source_scorecard": [],
        "source_scorecard_totals": dict(empty_scorecard_totals),
        "hub_performance": [],
        "incharge_scorecard": [],
        "incharge_scorecard_totals": dict(empty_scorecard_totals),
        "risk": [],
        "loss_register": [],
        "loss_register_total": 0,
        "loss_register_start": 0,
    }


def _resolve_snapshot_dates(requested_end):
    """(snapshot_date, latest_date). Default is the latest AVAILABLE BRSNR date;
    a requested date resolves to the last snapshot on/before it."""
    latest = frappe.db.sql(f"SELECT MAX(brsnr_date) FROM `tab{HISTORY_DOCTYPE}`")
    latest = latest[0][0] if latest and latest[0][0] else None

    if not latest or not requested_end:
        return latest, latest

    row = frappe.db.sql(
        f"SELECT MAX(brsnr_date) FROM `tab{HISTORY_DOCTYPE}` WHERE brsnr_date <= %s",
        (requested_end,),
    )
    return (row[0][0] if row and row[0][0] else None), latest


@frappe.whitelist(allow_guest=True)
def get_dashboard_data(
    email=None,
    brsnr_date=None,
    from_date=None,
    to_date=None,
    hub=None,
    current_hub=None,
    status=None,
    clearance_status=None,
    risk_level=None,
    search=None,
    limit=500,
    start=0,
    loss_limit=LOSS_REGISTER_LIMIT,
    loss_start=0,
    sort_by=None,
    sort_order=None,
    token=None,
):
    """BRSNR dashboard API with backend-enforced authorization.

    Authentication is the session token ONLY (X-BRSNR-Token / Authorization: Bearer /
    `token`). The `email` argument is accepted for backward compatibility and ignored.

    Incharge -> assigned Source hubs valid on the snapshot date.
    Director -> ALL hubs, including inactive ones.

    Semantics: snapshot/point-in-time data comes from Shipment History; New/Cleared
    are event counts over [from, snapshot]; Loss is cumulative THROUGH the selected
    date (loss_date <= date) from the master lifecycle, so the Loss Register
    survives after rows leave the Excel.
    """
    identity = require_portal_identity(token)

    status_values = _clean_choices(status, BRSNR_STATUSES, "status")
    clearance_values = _clean_choices(
        clearance_status,
        CLEARANCE_STATUSES,
        "clearance status",
    )
    risk_values = _clean_choices(risk_level, RISK_LEVELS, "risk level")

    from_req = _parse_date(from_date, "from date")
    to_req = _parse_date(to_date, "to date")
    if from_req and to_req and from_req > to_req:
        from_req, to_req = to_req, from_req

    requested_end = _parse_date(brsnr_date, "BRSNR date") or to_req
    snapshot_date, latest_date = _resolve_snapshot_dates(requested_end)

    if not snapshot_date:
        return _empty_dashboard(identity, [], None, latest_date)

    authorized, assignments, options = get_authorized_hubs(identity, snapshot_date)
    if not authorized:
        return _empty_dashboard(identity, [], snapshot_date, latest_date)

    # The hub filter can never escape the server-side scope.
    selected = authorized
    requested_hubs = _normalize_filter_list(hub)
    if requested_hubs and any(item.strip().lower() not in ("", "all") for item in requested_hubs):
        authorized_lookup = {str(h).strip().lower(): h for h in authorized}
        requested = []
        for item in requested_hubs:
            key = str(item).strip().lower()
            if not key or key == "all":
                continue
            match = authorized_lookup.get(key)
            if match and match not in requested:
                requested.append(match)
        selected = requested
        if not selected:
            frappe.throw(_("You are not authorized for the selected hub."), frappe.PermissionError)

    snapshot_date = getdate(snapshot_date)
    range_to = snapshot_date
    range_from = min(from_req or snapshot_date, range_to)
    trend_from = range_from if from_req else snapshot_date - timedelta(days=DEFAULT_TREND_DAYS - 1)
    trend_from = max(trend_from, snapshot_date - timedelta(days=MAX_TREND_DAYS - 1))

    params = {
        "hubs": tuple(selected),
        "d": snapshot_date,
        "rf": range_from,
        "rt": range_to,
        "tf": trend_from,
    }
    if status_values:
        params["status"] = tuple(status_values)
    if clearance_values:
        params["clearance"] = tuple(clearance_values)
    if risk_values:
        params["risk"] = tuple(risk_values)

    status_filter_sql = " AND x.brsnr_status IN %(status)s" if status_values else ""
    clearance_filter_sql = " AND x.clearance_status IN %(clearance)s" if clearance_values else ""
    risk_filter_sql = " AND x.risk_level IN %(risk)s" if risk_values else ""
    current_hub_values = _normalize_filter_list(current_hub)
    if current_hub_values:
        params["current_hubs"] = tuple(current_hub_values)
    current_hub_filter_sql = (
        " AND COALESCE(NULLIF(x.current_hub, ''), '(Unassigned)') IN %(current_hubs)s"
        if current_hub_values else ""
    )

    search_term = str(search or "").strip()[:100]
    if search_term:
        escaped = search_term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        params["q"] = f"%{escaped}%"
        search_filter_sql = (
            " AND (x.shipment_id LIKE %(q)s OR x.source LIKE %(q)s "
            "OR x.current_hub LIKE %(q)s OR s.casper_id LIKE %(q)s "
            "OR s.partner LIKE %(q)s)"
        )
    else:
        search_filter_sql = ""

    snapshot_filter_sql = (
        f"{status_filter_sql}{clearance_filter_sql}"
        f"{risk_filter_sql}{current_hub_filter_sql}{search_filter_sql}"
    )

    loss_filters = []
    if status_values:
        loss_filters.append("l.brsnr_status IN %(status)s")
    if clearance_values and "Loss" not in clearance_values:
        loss_filters.append("1 = 0")
    if risk_values and "Loss" not in risk_values:
        loss_filters.append("1 = 0")
    if current_hub_values:
        loss_filters.append(
            "COALESCE(NULLIF(l.current_hub, ''), '(Unassigned)') IN %(current_hubs)s"
        )
    if search_term:
        loss_filters.append(
            "(l.shipment_id LIKE %(q)s OR l.source LIKE %(q)s "
            "OR l.current_hub LIKE %(q)s OR l.partner LIKE %(q)s "
            "OR l.casper_id LIKE %(q)s)"
        )
    loss_filter_sql = (
        " AND " + " AND ".join(loss_filters) if loss_filters else ""
    )

    # ---- summary ---------------------------------------------------------
    snap = _sql(
        f"""
        SELECT
            COUNT(*) AS total,
            COALESCE(SUM(CASE WHEN x.brsnr_status = 'Pending' THEN 1 ELSE 0 END), 0) AS pending_count,
            COALESCE(SUM(CASE WHEN x.brsnr_status IN ('New', 'Pending') THEN 1 ELSE 0 END), 0) AS open_count,
            COALESCE(SUM(CASE WHEN x.clearance_status = 'Due Soon' THEN 1 ELSE 0 END), 0) AS due_soon,
            COALESCE(AVG(x.aging), 0) AS avg_aging
        FROM {_SNAP_ONE}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE 1 = 1 {snapshot_filter_sql}
        """,
        params,
    )
    snap = snap[0] if snap else {}

    events = {
        r.brsnr_status: cint(r.c)
        for r in _sql(
            f"""
            SELECT x.brsnr_status, COUNT(DISTINCT x.shipment, x.brsnr_date) AS c
            FROM {_SNAP_TREND}
            LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
            WHERE 1 = 1 {snapshot_filter_sql}
            GROUP BY x.brsnr_status
            """,
            params,
        )
    }

    loss_summary_rows = _sql(
        f"""
        SELECT COUNT(*) AS loss_count,
            COALESCE(SUM(l.loss_amount), 0) AS loss_value
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        """,
        params,
    )
    loss_summary = loss_summary_rows[0] if loss_summary_rows else {}
    loss_count = cint(loss_summary.get("loss_count"))
    loss_value = flt(loss_summary.get("loss_value"))

    risk_rows = _sql(
        f"""
        SELECT x.risk_level, COUNT(DISTINCT x.shipment) AS record_count,
            COALESCE(SUM(COALESCE(x.total_price, 0)), 0) AS amount
        FROM {_SNAP_ONE}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE x.risk_level IN ('Watch', 'Critical')
        AND 1 = 1 {snapshot_filter_sql}
        GROUP BY x.risk_level
        """,
        params,
    )
    risk_summary = {
        level: {"count": 0, "amount": 0.0}
        for level in ("Watch", "Critical", "Loss")
    }
    for row in risk_rows:
        risk_summary[row.risk_level] = {
            "count": cint(row.record_count),
            "amount": round(flt(row.amount), 2),
        }
    risk_summary["Loss"] = {
        "count": loss_count,
        "amount": round(loss_value, 2),
    }
    loss_events = _sql(
        f"""
        SELECT l.loss_date, COUNT(*) AS c,
            COALESCE(SUM(l.loss_amount), 0) AS v
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        GROUP BY l.loss_date
        ORDER BY l.loss_date ASC
        """,
        params,
    )

    cleared_range = events.get("Cleared", 0)
    open_count = cint(snap.get("open_count"))
    denominator = cleared_range + open_count
    avg_aging = round(flt(snap.get("avg_aging")), 2)

    summary = {
        "total": cint(snap.get("total")),
        "new": events.get("New", 0),
        "pending": cint(snap.get("pending_count")),
        "cleared": cleared_range,
        "clearance_percentage": round(cleared_range / denominator * 100, 2) if denominator else 0,
        "due_soon": cint(snap.get("due_soon")),
        "loss": loss_count,
        "loss_value": round(loss_value, 2),
        "avg_aging": avg_aging,
        "average_aging": avg_aging,
        "watch": risk_summary["Watch"]["count"],
        "watch_value": risk_summary["Watch"]["amount"],
        "critical": risk_summary["Critical"]["count"],
        "critical_value": risk_summary["Critical"]["amount"],
    }

    # ---- records ---------------------------------------------------------
    sort_expr = SORTABLE_FIELDS.get(str(sort_by or ""), "x.aging")
    direction = "ASC" if str(sort_order or "").lower() == "asc" else "DESC"

    page_limit = min(max(cint(limit or 500), 1), MAX_PAGE_LIMIT)
    page_start = max(cint(start or 0), 0)
    page_params = {**params, "lim": page_limit, "off": page_start}

    records = _sql(
        f"""
        SELECT
            x.shipment AS name, x.shipment_id, s.casper_id, s.profile_id, s.partner,
            s.wm_name, x.source, s.state, s.source_type, s.source_zone, s.rtorvpfwd,
            s.forward_reverse_type, s.`type` AS `type`, s.reverse_shipment_id,
            x.total_price, s.price_range, s.value_bucket, s.gmv, s.delivery_pincode,
            s.delivery_hub, x.current_hub, x.status, x.final_remarks, x.brsnr_status,
            x.brsnr_date, s.date_from, s.customer_promise_date, s.logistics_promise_date,
            s.latest_update_time, x.aging, x.aging_bucket, s.first_receive_time,
            s.first_receive_hub, s.last_receive_time, x.previous_status, x.previous_hub,
            IF(x.brsnr_status = 'Cleared', x.brsnr_date, NULL) AS cleared_date,
            x.brsnr_month, x.due_month, x.clearance_status,
            CASE WHEN x.clearance_status = 'Loss'
                THEN COALESCE(x.master_loss_date, x.due_date) END AS loss_date,
            x.loss_amount, x.days_remaining, x.risk_level
        FROM {_SNAP_ONE}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE 1 = 1 {snapshot_filter_sql}
        ORDER BY {sort_expr} {direction}, x.shipment_id ASC
        LIMIT %(lim)s OFFSET %(off)s
        """,
        page_params,
    )

    records_total = _sql(
        f"""
        SELECT COUNT(*) AS c
        FROM {_SNAP_ONE}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE 1 = 1 {snapshot_filter_sql}
        """,
        params,
    )
    records_total = cint(records_total[0].c) if records_total else 0

    # ---- trend (history based, spans the range; Loss is cumulative) -------
    trend = _sql(
        f"""
        SELECT
            x.brsnr_date,
            COUNT(DISTINCT x.shipment) AS total,
            COALESCE(SUM(CASE WHEN x.brsnr_status = 'New' THEN 1 ELSE 0 END), 0) AS new_count,
            COALESCE(SUM(CASE WHEN x.brsnr_status = 'Pending' THEN 1 ELSE 0 END), 0) AS pending_count,
            COALESCE(SUM(CASE WHEN x.brsnr_status = 'Cleared' THEN 1 ELSE 0 END), 0) AS cleared_count,
            COALESCE(SUM(CASE WHEN x.clearance_status = 'Due Soon' THEN 1 ELSE 0 END), 0) AS due_soon
        FROM {_SNAP_TREND}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE 1 = 1 {snapshot_filter_sql}
        GROUP BY x.brsnr_date
        ORDER BY x.brsnr_date ASC
        """,
        params,
    )
    pointer = 0
    running_loss_count = 0
    running_loss_amount = 0.0
    for point in trend:
        point_date = getdate(point.brsnr_date)
        while (
            pointer < len(loss_events)
            and getdate(loss_events[pointer].loss_date) <= point_date
        ):
            running_loss_count += cint(loss_events[pointer].c)
            running_loss_amount += flt(loss_events[pointer].v)
            pointer += 1
        point["loss_count"] = running_loss_count
        point["loss_value"] = round(running_loss_amount, 2)

    # ---- scorecards ------------------------------------------------------
    def _score_query(group_expression):
        return _sql(
            f"""
            SELECT {group_expression} AS hub,
                COUNT(*) AS total,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'New' THEN 1 ELSE 0 END), 0) AS new_count,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'Pending' THEN 1 ELSE 0 END), 0) AS pending_count,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'Cleared' THEN 1 ELSE 0 END), 0) AS cleared_count,
                COALESCE(SUM(CASE WHEN x.clearance_status = 'Due Soon' THEN 1 ELSE 0 END), 0) AS due_soon_count,
                COALESCE(SUM(CASE WHEN x.risk_level = 'Watch' THEN 1 ELSE 0 END), 0) AS watch_count,
                COALESCE(SUM(CASE WHEN x.risk_level = 'Critical' THEN 1 ELSE 0 END), 0) AS critical_count,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'New' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS new_amount,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'Pending' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS pending_amount,
                COALESCE(SUM(CASE WHEN x.brsnr_status = 'Cleared' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS cleared_amount,
                COALESCE(SUM(CASE WHEN x.clearance_status = 'Due Soon' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS due_soon_amount,
                COALESCE(SUM(CASE WHEN x.risk_level = 'Watch' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS watch_amount,
                COALESCE(SUM(CASE WHEN x.risk_level = 'Critical' THEN COALESCE(x.total_price, 0) ELSE 0 END), 0) AS critical_amount,
                COALESCE(SUM(COALESCE(x.total_price, 0)), 0) AS total_amount,
                COALESCE(AVG(x.aging), 0) AS avg_aging
            FROM {_SNAP_ONE}
            LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
            WHERE 1 = 1 {snapshot_filter_sql}
            GROUP BY {group_expression}
            """,
            params,
        )

    loss_by_current_hub = _sql(
        f"""
        SELECT COALESCE(NULLIF(l.current_hub, ''), '(Unassigned)') AS hub,
            COUNT(*) AS loss_count, COALESCE(SUM(l.loss_amount), 0) AS loss_amount
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        GROUP BY COALESCE(NULLIF(l.current_hub, ''), '(Unassigned)')
        """,
        params,
    )
    loss_by_source = _sql(
        f"""
        SELECT l.source AS hub, COUNT(*) AS loss_count,
            COALESCE(SUM(l.loss_amount), 0) AS loss_amount
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        GROUP BY l.source
        """,
        params,
    )

    hub_performance = {}
    score_fields = (
        "total", "new_count", "pending_count", "cleared_count",
        "due_soon_count", "watch_count", "critical_count",
        "new_amount", "pending_amount", "cleared_amount", "due_soon_amount",
        "watch_amount", "critical_amount", "total_amount", "avg_aging",
    )
    for row in _score_query("COALESCE(NULLIF(x.current_hub, ''), '(Unassigned)')"):
        hub_performance[row.hub] = {
            "hub": row.hub,
            "hub_name": row.hub,
            **{field: round(flt(row.get(field)), 2) if field.endswith("amount") or field == "avg_aging"
               else cint(row.get(field)) for field in score_fields},
            "loss_count": 0,
            "loss_amount": 0.0,
        }
    for row in loss_by_current_hub:
        score = hub_performance.setdefault(
            row.hub,
            {
                "hub": row.hub, "hub_name": row.hub,
                **{field: 0 for field in score_fields},
                "loss_count": 0, "loss_amount": 0.0,
            },
        )
        score["loss_count"] = cint(row.loss_count)
        score["loss_amount"] = round(flt(row.loss_amount), 2)
    for score in hub_performance.values():
        score["loss_value"] = score["loss_amount"]
        denominator = score["cleared_count"] + score["pending_count"]
        score["clearance_percentage"] = (
            round(score["cleared_count"] / denominator * 100, 2) if denominator else 0
        )
    hub_performance = sorted(
        hub_performance.values(),
        key=lambda row: (-row["total"], str(row["hub"]).lower()),
    )

    scorecard_totals = {
        field: round(sum(flt(row[field]) for row in hub_performance), 2)
        if field.endswith("amount")
        else sum(cint(row[field]) for row in hub_performance)
        for field in score_fields + ("loss_count", "loss_amount")
    }
    scorecard_totals["avg_aging"] = round(
        sum(flt(row["avg_aging"]) * cint(row["total"]) for row in hub_performance)
        / scorecard_totals["total"],
        2,
    ) if scorecard_totals["total"] else 0
    scorecard_totals["loss_value"] = scorecard_totals["loss_amount"]

    source_metrics = {
        row.hub: {field: row.get(field) for field in score_fields}
        for row in _score_query("x.source")
    }
    source_losses = {
        row.hub: {
            "loss_count": cint(row.loss_count),
            "loss_amount": round(flt(row.loss_amount), 2),
        }
        for row in loss_by_source
    }
    source_names = {row["name"]: row.get("hub_name") for row in options}
    source_performance = {}
    for source, metrics in source_metrics.items():
        source_performance[source] = {
            "hub": source,
            "hub_name": source_names.get(source),
            **{
                field: round(flt(metrics.get(field)), 2)
                if field.endswith("amount") or field == "avg_aging"
                else cint(metrics.get(field))
                for field in score_fields
            },
            "loss_count": 0,
            "loss_amount": 0.0,
            "loss_value": 0.0,
        }
    for source, loss in source_losses.items():
        score = source_performance.setdefault(
            source,
            {
                "hub": source,
                "hub_name": source_names.get(source),
                "total": 0,
                "new_count": 0,
                "pending_count": 0,
                "cleared_count": 0,
                "loss_count": 0,
                "loss_amount": 0.0,
                "loss_value": 0.0,
                "avg_aging": 0,
            },
        )
        score["loss_count"] = loss["loss_count"]
        score["loss_amount"] = loss["loss_amount"]
        score["loss_value"] = loss["loss_amount"]
    for score in source_performance.values():
        denominator = score["cleared_count"] + score["pending_count"]
        score["clearance_percentage"] = (
            round(score["cleared_count"] / denominator * 100, 2) if denominator else 0
        )
    source_performance = sorted(
        source_performance.values(),
        key=lambda row: (-row["pending_count"], -row["loss_count"]),
    )

    source_scorecard_map = {}
    for hub_option in options:
        hub_name = hub_option["name"]
        source_scorecard_map[hub_name] = {
            "hub": hub_name,
            "source_hub": hub_name,
            "hub_name": hub_option.get("hub_name") or hub_name,
            **{field: 0 for field in score_fields},
            "loss_count": 0,
            "loss_amount": 0.0,
            "loss_value": 0.0,
        }
    for row in source_performance:
        score = source_scorecard_map.get(row["hub"])
        if score:
            score.update(row)
            score["source_hub"] = row["hub"]
    source_scorecard = sorted(
        source_scorecard_map.values(),
        key=lambda row: str(row["hub_name"]).lower(),
    )
    source_scorecard_totals = {
        field: round(sum(flt(row[field]) for row in source_scorecard), 2)
        if field.endswith("amount")
        else sum(cint(row[field]) for row in source_scorecard)
        for field in score_fields + ("loss_count", "loss_amount")
    }
    source_scorecard_totals["avg_aging"] = round(
        sum(flt(row["avg_aging"]) * cint(row["total"]) for row in source_scorecard)
        / source_scorecard_totals["total"],
        2,
    ) if source_scorecard_totals["total"] else 0
    source_scorecard_totals["loss_value"] = source_scorecard_totals["loss_amount"]

    assignment_where = [
        "a.active = 1",
        "(a.from_date IS NULL OR a.from_date <= %(d)s)",
        "(a.to_date IS NULL OR a.to_date >= %(d)s)",
        "i.active = 1",
        "a.hub IN %(hubs)s",
    ]
    if identity["type"] == "Incharge":
        assignment_where.append("a.hub_incharge = %(incharge)s")
        params["incharge"] = identity["name"]
    assignment_rows = _sql(
        f"""
        SELECT DISTINCT a.hub_incharge, i.incharge_name, a.hub,
            h.hub_name, h.hub_code
        FROM `tab{ASSIGNMENT_DOCTYPE}` a
        INNER JOIN `tab{INCHARGE_DOCTYPE}` i ON i.name = a.hub_incharge
        INNER JOIN `tab{HUB_DOCTYPE}` h ON h.name = a.hub
        WHERE {" AND ".join(assignment_where)}
        ORDER BY i.incharge_name, a.hub
        """,
        params,
    )
    assignment_owner = {}
    incharge_data = {}
    for row in assignment_rows:
        # A hub may be assigned to more than one active incharge on the same
        # snapshot date. In that case, aggregate the same hub under each
        # incharge instead of failing the dashboard request.
        assignment_owner.setdefault(row.hub, set()).add(row.hub_incharge)

        item = incharge_data.setdefault(
            row.hub_incharge,
            {
                "incharge": row.hub_incharge,
                "incharge_name": row.incharge_name,
                "assigned_hubs": [],
                "hubs": [],
                **{field: 0 for field in score_fields + ("loss_count", "loss_amount")},
                "aging_weighted": 0.0,
            },
        )
        if row.hub not in item["assigned_hubs"]:
            item["assigned_hubs"].append(row.hub)
            source = source_scorecard_map.get(row.hub)
            if not source:
                continue
            item["hubs"].append(dict(source))
            for field in score_fields:
                if field != "avg_aging":
                    item[field] += flt(source[field])
            item["aging_weighted"] += (
                flt(source["avg_aging"]) * cint(source["total"])
            )
            item["loss_count"] += cint(source["loss_count"])
            item["loss_amount"] += flt(source["loss_amount"])

    unassigned_hubs = [
        dict(row) for row in source_scorecard_map.values()
        if row["hub"] not in assignment_owner
    ]
    if unassigned_hubs:
        item = incharge_data.setdefault(
            None,
            {
                "incharge": None,
                "incharge_name": _("Unassigned"),
                "assigned_hubs": [],
                "hubs": [],
                **{field: 0 for field in score_fields + ("loss_count", "loss_amount")},
                "aging_weighted": 0.0,
            },
        )
        for source in unassigned_hubs:
            item["assigned_hubs"].append(source["hub"])
            item["hubs"].append(source)
            for field in score_fields:
                if field != "avg_aging":
                    item[field] += flt(source[field])
            item["aging_weighted"] += flt(source["avg_aging"]) * cint(source["total"])
            item["loss_count"] += cint(source["loss_count"])
            item["loss_amount"] += flt(source["loss_amount"])

    incharge_scorecard = []
    for item in incharge_data.values():
        item["assigned_hubs"].sort()
        item["hubs"].sort(key=lambda row: str(row["hub_name"]).lower())
        item["assigned_hubs_text"] = ", ".join(item["assigned_hubs"])
        item["loss_value"] = round(item["loss_amount"], 2)
        item["loss_amount"] = round(item["loss_amount"], 2)
        item["avg_aging"] = round(
            item["aging_weighted"] / item["total"], 2
        ) if item["total"] else 0
        del item["aging_weighted"]
        for field in score_fields:
            if field.endswith("amount"):
                item[field] = round(item[field], 2)
        for hub_row in item["hubs"]:
            hub_row["loss_value"] = hub_row["loss_amount"]
        if not search_term or item["total"] or search_term.lower() in (
            item["incharge_name"] + " " + item["assigned_hubs_text"]
        ).lower():
            incharge_scorecard.append(item)
    incharge_scorecard.sort(
        key=lambda row: (row["incharge"] is None, row["incharge_name"].lower())
    )
    incharge_scorecard_totals = {
        field: source_scorecard_totals[field]
        for field in score_fields + ("loss_count", "loss_amount")
    }
    incharge_scorecard_totals["avg_aging"] = source_scorecard_totals["avg_aging"]
    incharge_scorecard_totals["loss_value"] = source_scorecard_totals["loss_amount"]

    # ---- risk and Loss Register -----------------------------------------
    risk = _sql(
        f"""
        SELECT x.shipment AS name, x.shipment_id, x.source, x.current_hub,
            x.total_price, x.aging, x.brsnr_month, x.due_month,
            x.clearance_status, x.loss_amount, x.days_remaining, x.risk_level
        FROM {_SNAP_ONE}
        LEFT JOIN `tab{SHIPMENT_DOCTYPE}` s ON s.name = x.shipment
        WHERE x.risk_level IN ('Watch', 'Critical')
        AND 1 = 1 {snapshot_filter_sql}
        ORDER BY CASE WHEN x.risk_level = 'Critical' THEN 0 ELSE 1 END,
            x.aging DESC, x.shipment_id ASC
        LIMIT 100
        """,
        params,
    )
    loss_total_rows = _sql(
        f"""
        SELECT COUNT(*) AS c
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        """,
        params,
    )
    loss_register_total = cint(loss_total_rows[0].c) if loss_total_rows else 0
    loss_page_limit = min(max(cint(loss_limit or LOSS_REGISTER_LIMIT), 1), MAX_PAGE_LIMIT)
    loss_page_start = max(cint(loss_start or 0), 0)
    loss_page_params = {
        **params,
        "loss_lim": loss_page_limit,
        "loss_off": loss_page_start,
    }
    loss_sortable = {
        "loss_date": "l.loss_date",
        "loss_amount": "l.loss_amount",
        "total_price": "l.total_price",
        "shipment_id": "l.shipment_id",
        "source": "l.source",
        "current_hub": "l.current_hub",
        "brsnr_status": "l.brsnr_status",
    }
    loss_sort_expr = loss_sortable.get(str(sort_by or ""), "l.loss_date")
    loss_sort_direction = "ASC" if str(sort_order or "").lower() == "asc" else "DESC"
    loss_register = _sql(
        f"""
        SELECT l.*
        FROM {_loss_register_sql()}
        WHERE 1 = 1 {loss_filter_sql}
        ORDER BY {loss_sort_expr} {loss_sort_direction}, l.shipment_id ASC
        LIMIT %(loss_lim)s OFFSET %(loss_off)s
        """,
        loss_page_params,
    )

    return {
        "success": True,
        "role": identity["type"],
        "incharge": identity["display_name"],
        "email": identity["email"],
        "hubs": selected,
        "all_authorized_hubs": authorized,
        "hub_options": options,
        "assignments": assignments if identity["type"] == "Incharge" else [],
        "brsnr_date": snapshot_date,
        "latest_brsnr_date": latest_date,
        "filters": {
            "from_date": range_from,
            "to_date": range_to,
            "hub": hub,
            "current_hub": current_hub,
            "status": status,
            "clearance_status": clearance_status,
            "risk_level": risk_level,
            "search": search,
        },
        "summary": summary,
        "risk_summary": risk_summary,
        "records": records,
        "records_total": records_total,
        "trend": trend,
        "hub_scorecard": hub_performance,
        "hub_scorecard_totals": scorecard_totals,
        "source_scorecard": source_scorecard,
        "source_scorecard_totals": source_scorecard_totals,
        "hub_performance": source_performance,
        "incharge_scorecard": incharge_scorecard,
        "incharge_scorecard_totals": incharge_scorecard_totals,
        "risk": risk,
        "loss_register": loss_register,
        "loss_register_total": loss_register_total,
        "loss_register_start": loss_page_start,
    }