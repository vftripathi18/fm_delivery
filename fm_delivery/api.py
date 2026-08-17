import frappe

from fm_delivery.utils.excel_import import ExcelImporter

from openpyxl import load_workbook


# ============================================================
# DAILY DELIVERY IMPORT
# ============================================================

@frappe.whitelist()
def import_excel(docname):

    import_doc = frappe.get_doc(
        "FM Delivery Import",
        docname
    )

    importer = ExcelImporter(import_doc)

    importer.load_workbook()

    importer.validate_headers()

    result = importer.import_data()

    import_doc.status = "Completed"

    import_doc.save(
        ignore_permissions=True
    )

    return result


# ============================================================
# SALARY DECLARATION IMPORT
# ============================================================

@frappe.whitelist()
def import_salary_declaration(docname):

    # --------------------------------------------------------
    # Get Parent Document
    # --------------------------------------------------------

    doc = frappe.get_doc(
        "FM Salary Declaration",
        docname
    )

    # --------------------------------------------------------
    # Validate Dates
    # --------------------------------------------------------

    if not doc.from_date:
        frappe.throw(
            "From Date is required."
        )

    if not doc.to_date:
        frappe.throw(
            "To Date is required."
        )

    if doc.from_date > doc.to_date:
        frappe.throw(
            "From Date cannot be greater than To Date."
        )

    # --------------------------------------------------------
    # Validate Excel
    # --------------------------------------------------------

    if not doc.salary_excel:
        frappe.throw(
            "Please attach the salary Excel file."
        )

    # --------------------------------------------------------
    # Get Excel File
    # --------------------------------------------------------

    file_doc = frappe.get_doc(
        "File",
        {
            "file_url": doc.salary_excel
        }
    )

    file_path = file_doc.get_full_path()

    # --------------------------------------------------------
    # Load Excel
    # --------------------------------------------------------

    try:

        workbook = load_workbook(
            filename=file_path,
            data_only=True,
            read_only=True
        )

    except Exception as e:

        frappe.throw(
            f"Unable to read Excel file: {str(e)}"
        )

    sheet = workbook.active

    # --------------------------------------------------------
    # Expected Excel Headers
    # --------------------------------------------------------

    expected_headers = [
        "State",
        "Hub_Name",
        "FIX Salary",
        "FWD Rate",
        "RTO Rate",
        "Total Rate",
    ]

    # --------------------------------------------------------
    # Read Headers
    # --------------------------------------------------------

    headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in sheet[1]
    ]

    # --------------------------------------------------------
    # Validate Headers
    # --------------------------------------------------------

    if headers != expected_headers:

        frappe.throw(
            f"""
            <h4>Invalid Excel Headers</h4>

            <b>Expected:</b>
            <br>
            {'<br>'.join(expected_headers)}

            <hr>

            <b>Found:</b>
            <br>
            {'<br>'.join(headers)}
            """
        )

    # --------------------------------------------------------
    # Read Excel Rows
    # --------------------------------------------------------

    rows = list(
        sheet.iter_rows(
            min_row=2,
            values_only=True
        )
    )

    # Remove completely blank rows

    rows = [
        row
        for row in rows
        if any(
            value not in (None, "")
            for value in row
        )
    ]

    total_rows = len(rows)

    # --------------------------------------------------------
    # Clear Existing Child Table
    # --------------------------------------------------------

    doc.set(
        "salary_details",
        []
    )

    imported = 0
    failed = 0

    # --------------------------------------------------------
    # Process Rows
    # --------------------------------------------------------

    for row_number, row in enumerate(
        rows,
        start=2
    ):

        try:

            # ------------------------------------------------
            # Read Values
            # ------------------------------------------------

            state = str(
                row[0] or ""
            ).strip()

            hub_name = str(
                row[1] or ""
            ).strip()

            fix_salary = (
                float(row[2])
                if row[2] not in (None, "")
                else 0
            )

            fwd_rate = (
                float(row[3])
                if row[3] not in (None, "")
                else 0
            )

            rto_rate = (
                float(row[4])
                if row[4] not in (None, "")
                else 0
            )

            total_rate = (
                float(row[5])
                if row[5] not in (None, "")
                else 0
            )

            # ------------------------------------------------
            # Required Validation
            # ------------------------------------------------

            if not state:

                frappe.throw(
                    f"Row {row_number}: State is required."
                )

            if not hub_name:

                frappe.throw(
                    f"Row {row_number}: Hub_Name is required."
                )

            # ------------------------------------------------
            # Salary Method Validation
            # ------------------------------------------------

            methods = 0

            # FIX Salary method
            if fix_salary > 0:
                methods += 1

            # FWD/RTO method
            if fwd_rate > 0 or rto_rate > 0:
                methods += 1

            # Total Rate method
            if total_rate > 0:
                methods += 1

            # No method
            if methods == 0:

                frappe.throw(
                    f"""
                    Row {row_number}: No salary configuration found
                    for:

                    State: {state}
                    Hub: {hub_name}

                    Please provide FIX Salary,
                    FWD/RTO Rate, or Total Rate.
                    """
                )

            # Multiple methods
            if methods > 1:

                frappe.throw(
                    f"""
                    Row {row_number}: Multiple salary configurations
                    found for:

                    State: {state}
                    Hub: {hub_name}

                    Only one salary method is allowed:

                    1. FIX Salary
                    OR
                    2. FWD/RTO Rate
                    OR
                    3. Total Rate
                    """
                )

            # ------------------------------------------------
            # Add Child Row
            # ------------------------------------------------

            child = doc.append(
                "salary_details",
                {}
            )

            child.state = state

            child.hub_name = hub_name

            child.fix_salary = fix_salary

            child.fwd_rate = fwd_rate

            child.rto_rate = rto_rate

            child.total_rate = total_rate

            imported += 1

        except Exception as e:

            failed += 1

            frappe.log_error(
                title=f"Salary Declaration Import Row {row_number}",
                message=frappe.get_traceback()
            )

            frappe.throw(
                f"Error in Excel row {row_number}: {str(e)}"
            )

    # --------------------------------------------------------
    # Save Parent
    # --------------------------------------------------------

    doc.save(
        ignore_permissions=True
    )

    frappe.db.commit()

    # --------------------------------------------------------
    # Return Result
    # --------------------------------------------------------

    return {
        "total_rows": total_rows,
        "imported": imported,
        "failed": failed,
    }